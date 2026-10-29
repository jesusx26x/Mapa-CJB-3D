"""
Exporta el modelo «Flujo de Caja VBC Oct2026-Dic2027 v*.xlsx» (desde v1.11, «…Oct2026-Dic2028 v*.xlsx»)
a data/flujo_2027.json para la pestaña «💰 Flujo 2027» del Sistema VBC.

Sin --src usa la versión más alta del modelo en la carpeta 2027 (v1, v1.1, v1.2…;
a igual versión, el archivo guardado más recientemente). Los meses salen del Excel
(Flujo_Mensual, fila 4): 15 hasta v1.10 (oct-26 a dic-27), 27 desde v1.11 (hasta dic-28).

Desde v1.19 el Panel elige, además del escenario (D5), la versión de ventas y obras (D7, nombre VersionActiva;
lista en Decisiones, sección 8). Se exportan todas las combinaciones versión × escenario (8 × 3 = 24), cada una
recalculada por Excel, para que la página cambie de versión y de escenario sin volver a exportar.

Pasos:
  1. Copia el Excel a una carpeta temporal (el original nunca se abre ni se guarda).
  2. Por cada versión y escenario escribe la versión en Panel!D7 y el escenario en Panel!D5, recalcula con Excel
     por COM (PowerShell + Excel.Application, sin dependencias nuevas) y guarda una copia por combinación.
  3. Lee las copias con openpyxl (data_only=True), validando los encabezados de cada hoja.
  4. Cuadra los detalles contra Flujo_Mensual y el Panel en cada combinación, compara la Base de cada versión con
     los valores fijos de la hoja Versiones y escribe el JSON (esquema 2: «combos» por «versión|escenario»).

Uso:  python tools/export_flujo_2027.py [--src RUTA_EXCEL] [--out RUTA_JSON] [--solo-activa]
      --solo-activa: solo la versión elegida en el Panel (3 recálculos en vez de 24).
"""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
MODELO_DIR = Path(
    r"C:\Users\John Marte\Fideicomiso VBCRD\Fonvivienda - IDMP\VBC2\Planificación Estratégica"
    r"\01 Autosuficiencia Financiera\Presupuesto\2027"
)
MODELO_PATRON = "Flujo de Caja VBC Oct2026-Dic20?? v*.xlsx"  # Dic2027 hasta v1.10, Dic2028 desde v1.11
DEFAULT_OUT = ROOT / "data" / "flujo_2027.json"
ESCENARIOS = ["Pesimista", "Base", "Optimista"]
TOL = 1.0  # RD$ de tolerancia al cuadrar

# Claves con lote real en el mapa: CJB-04-13-01, CPAG-01-01-01, CPAG-C-05
RE_CLAVE_REAL = re.compile(r"^(CJB|CPAG)-(\d{2}-\d{2}-\d{2}|C-\d{2})$")

POWERSHELL_RECALC = r"""
$ErrorActionPreference = 'Stop'
$src = $env:VBC_FLUJO_SRC
$dir = $env:VBC_FLUJO_DIR
$escenarios = $env:VBC_FLUJO_ESC -split ','
# Versiones a recalcular: «n|etiqueta de la lista del Panel» por línea (UTF-8). Sin archivo (modelos anteriores a
# v1.19, sin selector de versión): solo los escenarios, como versión 0.
$versiones = @()
if ($env:VBC_FLUJO_VER) { $versiones = @(Get-Content -LiteralPath $env:VBC_FLUJO_VER -Encoding UTF8 | Where-Object { $_ -ne '' }) }
if ($versiones.Count -eq 0) { $versiones = @('0|') }
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false
$xl.DisplayAlerts = $false
$xl.AskToUpdateLinks = $false
try {
  $wb = $xl.Workbooks.Open($src, 0, $false)
  $panel = $wb.Worksheets.Item('Panel')
  foreach ($v in $versiones) {
    $n, $etiqueta = $v -split '\|', 2
    if ($etiqueta) { $wb.Names.Item('VersionActiva').RefersToRange.Value2 = $etiqueta }
    foreach ($e in $escenarios) {
      $panel.Range('D5').Value2 = $e
      $xl.CalculateFull()
      $dest = Join-Path $dir ("flujo_v" + $n + "_" + $e + ".xlsx")
      $wb.SaveCopyAs($dest)
      Write-Output ("OK v" + $n + " " + $e)
    }
  }
  $wb.Close($false)
} finally {
  $xl.Quit()
  [void][Runtime.InteropServices.Marshal]::ReleaseComObject($xl)
}
"""


# ----------------------------------------------------------------------------- utilidades
def norm(v):
    """Normaliza un encabezado para compararlo (minúsculas, sin espacios extra)."""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip().lower()


def mes_key(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return f"{v.year:04d}-{v.month:02d}"
    return None


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return 0.0


def r2(x):
    return round(float(x), 2)


def fail(msg):
    raise SystemExit(f"ERROR: {msg}")


def find_header_row(ws, required, max_row=12):
    """Devuelve (fila, {encabezado_normalizado: columna}) de la primera fila que tenga todos los requeridos."""
    req = [norm(r) for r in required]
    for r in range(1, max_row + 1):
        hdr = {}
        for c in range(1, ws.max_column + 1):
            h = norm(ws.cell(r, c).value)
            if h and h not in hdr:
                hdr[h] = c
        if all(x in hdr for x in req):
            return r, hdr
    fail(f"Hoja «{ws.title}»: no encontré los encabezados {required}")


def month_cols(ws, row, meses):
    """Columnas del primer bloque contiguo de fechas en `row` que coincide con `meses`."""
    for c in range(1, ws.max_column + 1):
        if mes_key(ws.cell(row, c).value) == meses[0]:
            cols = list(range(c, c + len(meses)))
            got = [mes_key(ws.cell(row, cc).value) for cc in cols]
            if got == meses:
                return cols
    fail(f"Hoja «{ws.title}» fila {row}: no encontré los meses {meses[0]}…{meses[-1]} contiguos")


def etiqueta_posterior(ws, prefijo=""):
    """Encabezado de la columna «Posterior»: «(2028+)» hasta v1.10, «(2029+)» desde v1.11."""
    pat = re.compile("^" + (re.escape(norm(prefijo)) + " " if prefijo else "") + r"posterior \(20\d\d\+\)$")
    for r in range(1, 13):
        for c in range(1, ws.max_column + 1):
            if pat.match(norm(ws.cell(r, c).value)):
                return norm(ws.cell(r, c).value)
    fail(f"Hoja «{ws.title}»: no encontré la columna «{prefijo}Posterior (20xx+)»")


def meses_flujo(ws):
    """Meses de Flujo_Mensual (fila 4, desde la columna B hasta la primera celda que no es fecha)."""
    meses, c = [], 2
    while mes_key(ws.cell(4, c).value):
        meses.append(mes_key(ws.cell(4, c).value))
        c += 1
    if len(meses) < 15 or len(set(meses)) != len(meses):
        fail(f"Flujo_Mensual: la fila 4 no tiene meses consecutivos desde B: {meses}")
    return meses


def col(hdr, name, sheet):
    k = norm(name)
    if k not in hdr:
        fail(f"Hoja «{sheet}»: falta la columna «{name}»")
    return hdr[k]


def data_rows(ws, start, key_col):
    """Filas de datos desde `start` hasta la primera fila vacía o 'TOTAL' en key_col."""
    r = start
    while r <= ws.max_row:
        v = ws.cell(r, key_col).value
        if v is None or norm(v) == "total":
            break
        yield r
        r += 1


def table_rows(ws, name):
    """Filas de una tabla de Excel como lista de dicts por encabezado."""
    if name not in ws.tables:
        fail(f"Hoja «{ws.title}»: no existe la tabla «{name}»")
    ref = ws.tables[name].ref
    rows = list(ws[ref])
    heads = [c.value for c in rows[0]]
    out = []
    for row in rows[1:]:
        d = {h: c.value for h, c in zip(heads, row)}
        if any(v is not None for v in d.values()):
            out.append(d)
    return heads, out


def split_clave(clave):
    m = RE_CLAVE_REAL.match(clave or "")
    return (m.group(1), m.group(2)) if m else (None, None)


def zeros(n):
    return [0.0] * n


def add_into(acc, arr):
    for i, v in enumerate(arr):
        acc[i] += v


def celda_nombre(wb, nombre):
    """Celda de un nombre definido de una sola celda (p. ej. «VersionActiva» = Panel!D7); None si no existe."""
    dn = wb.defined_names.get(nombre)
    if dn is None:
        return None
    for hoja, coord in dn.destinations:
        if hoja in wb.sheetnames and ":" not in coord:
            return wb[hoja][coord.replace("$", "")]
    return None


def valor_nombre(wb, nombre, defecto=None):
    c = celda_nombre(wb, nombre)
    return defecto if c is None else c.value


def valor_celda(v, ratio=False):
    """Valor de una celda para el JSON: números redondeados (6 decimales si es un %), fechas como aaaa-mm, texto tal cual."""
    if isinstance(v, (dt.datetime, dt.date)):
        return mes_key(v)
    if isinstance(v, (int, float)):
        return round(float(v), 6) if ratio else r2(v)
    return "" if v is None else str(v).strip()


def filas_etiqueta(ws, col=2, r0=1, r1=None):
    """[(etiqueta normalizada, fila)] de las celdas no vacías de una columna, en orden de fila."""
    r1 = min(r1 or ws.max_row, ws.max_row)
    return [(norm(ws.cell(r, col).value), r) for r in range(r0, r1 + 1) if norm(ws.cell(r, col).value)]


def buscar_fila(filas, inicio, hoja, opcional=False, exacta=False):
    """Primera fila cuya etiqueta empieza por `inicio` (o es igual, con exacta=True). Las hojas de resumen se leen por
    etiqueta y no por número de fila, porque cada versión del modelo inserta filas."""
    p = norm(inicio)
    r = next((r for lab, r in filas if (lab == p if exacta else lab.startswith(p))), None)
    if r is None and not opcional:
        fail(f"Hoja «{hoja}»: no encontré la fila «{inicio}»")
    return r


def columnas_encabezado(ws, fila, campos, hoja, c0=2):
    """{clave: columna} buscando en la fila de encabezado cada texto inicial de `campos` [(clave, inicio)]."""
    hdr = [(norm(ws.cell(fila, c).value), c) for c in range(c0, ws.max_column + 1) if norm(ws.cell(fila, c).value)]
    out = {}
    for k, ini in campos:
        c = next((c for h, c in hdr if h.startswith(norm(ini))), None)
        if c is None:
            fail(f"Hoja «{hoja}» fila {fila}: falta la columna «{ini}…»")
        out[k] = c
    return out


def columnas_mes(ws, fila, c0=3):
    """[(aaaa-mm, columna)] de las fechas contiguas de una fila, desde la columna c0."""
    out, c = [], c0
    while mes_key(ws.cell(fila, c).value):
        out.append((mes_key(ws.cell(fila, c).value), c))
        c += 1
    return out


# ----------------------------------------------------------------------------- recálculo
def copiar_modelo(src, tmpdir):
    copia = Path(tmpdir) / "modelo.xlsx"
    try:
        shutil.copyfile(src, copia)  # el original no se toca
    except PermissionError:
        # Archivos de OneDrive (reparse points): Python no siempre puede abrirlos; Copy-Item sí.
        res = subprocess.run(
            ["powershell", "-NoProfile", "-Command",
             "Copy-Item -LiteralPath $env:VBC_FLUJO_ORIG -Destination $env:VBC_FLUJO_COPIA -Force"],
            env=dict(os.environ, VBC_FLUJO_ORIG=str(src), VBC_FLUJO_COPIA=str(copia)),
            capture_output=True, text=True, timeout=120,
        )
        if res.returncode != 0 or not copia.exists():
            fail("No pude copiar el Excel a la carpeta temporal:\n" + res.stdout + res.stderr)
    return copia


def recalcular(copia, tmpdir, versiones):
    """Recalcula cada versión × escenario. `versiones` = [(n, etiqueta del Panel D7)]; vacía en modelos sin selector de
    versión (antes de v1.19), que se recalculan solo por escenario como versión 0. Devuelve {(n, escenario): copia}."""
    env = dict(os.environ)
    env.update(VBC_FLUJO_SRC=str(copia), VBC_FLUJO_DIR=str(tmpdir), VBC_FLUJO_ESC=",".join(ESCENARIOS))
    if versiones:
        lista = Path(tmpdir) / "versiones.txt"
        lista.write_text("\n".join(f"{n}|{et}" for n, et in versiones), encoding="utf-8-sig")
        env["VBC_FLUJO_VER"] = str(lista)
    ps1 = Path(tmpdir) / "recalc.ps1"
    ps1.write_text(POWERSHELL_RECALC, encoding="utf-8-sig")
    nums = [n for n, _ in versiones] or [0]
    print(f"Recalculando con Excel (COM): {len(nums)} versión(es) × {len(ESCENARIOS)} escenarios…")
    res = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps1)],
        env=env, capture_output=True, text=True, timeout=1800,
    )
    if res.returncode != 0:
        fail("Excel COM falló:\n" + res.stdout + res.stderr)
    out = {}
    for n in nums:
        for e in ESCENARIOS:
            p = Path(tmpdir) / f"flujo_v{n}_{e}.xlsx"
            if not p.exists():
                fail(f"No se generó la copia recalculada de la versión {n}, escenario {e}")
            out[(n, e)] = p
    return out


# ----------------------------------------------------------------------------- lectura
# Indicador -> etiquetas de la columna B del Panel (v1.11 primero; luego la de v1.10 y anteriores)
PANEL_KPIS = [
    ("saldoInicial", ["Saldo inicial de caja"]),
    ("ingresos", ["Ingresos proyectados"]),
    ("egresos", ["Egresos proyectados"]),
    ("saldoFinal", ["Saldo final del período", "Saldo final a dic-2027"]),
    ("saldoMin", ["Saldo más bajo del período"]),
    ("mesSaldoMin", ["Mes del saldo más bajo"]),
    ("mesesNegativos", ["Meses con saldo negativo"]),
    ("colchon", ["Colchón mínimo de caja"]),
    ("disponiblePACC", ["Disponible para presupuesto de áreas (PACC)"]),
    ("disponibleCierre", ["Disponible al cierre del período", "Disponible al cierre de dic-2027"]),
]

FLUJO_FILAS = {
    "saldoInicial": "Saldo inicial",
    "cartera": "Cartera con plan de pagos",
    "recupLegal": "Recuperación de cartera en gestión legal",
    "lic2026": "Licitación 2026-01",
    "lic2027": "Licitación 2027-01",
    "zfi": "Zona Franca Industrial",
    "adicionales": "Lotes adicionales",
    "cpag2027": "Licitaciones CPAG 2027",
    "inversiones": "Ingresos por inversiones",
    "otros": "Otros ingresos",
    "totalIngresos": "TOTAL INGRESOS",
    "obrasCJB": "Cubicaciones obras CJB",
    "obrasCPAG": "Cubicaciones obras CPAG",
    "mantGeneral": ("Mantenimiento de la ciudad", "Reserva Mantenimiento General"),  # v1.14 / anteriores
    "complementarias": "Obras complementarias",
    "costoAtraso": "Costo de atrasar obras",
    "retenidos": "Devolución de retenidos",
    "gfRRHH": "Gastos fijos – Recursos Humanos",
    "gfCargaFija": "Gastos fijos – Carga Fija",
    "gfAdmin": "Gastos fijos – Administrativos",
    "gfMant": "Gastos fijos – Mantenimientos",
    "gfITBIS": "Gastos fijos – ITBIS",
    "gfProvision": "Provisión de gastos",
    "obligaciones": "Obligaciones adicionales",
    "gfBienes": ("Bienes y servicios de las áreas", "Bienes muebles"),  # v1.14: compras de las áreas (techos PACC)
    "totalEgresos": "TOTAL EGRESOS",
    "neto": "FLUJO NETO DEL MES",
    "lineaDesembolso": "Desembolso de línea de crédito",
    "lineaPago": "Pago de línea de crédito",
    "lineaIntereses": "Intereses de línea de crédito",
    "lineaDeuda": "Deuda de la línea de crédito",
    "saldo": "SALDO FINAL DE CAJA",
    "colchon": "Colchón mínimo",
    "estado": "Estado",
}
# Filas que solo existen desde v1.12 (gastos reales de Finanzas), v1.13 (inversiones) y v1.14 (obligaciones):
# en versiones anteriores valen 0.
FLUJO_FILAS_OPCIONALES = {"gfITBIS", "gfProvision", "inversiones", "obligaciones"}
# Las compras de las áreas (gfBienes) no son gasto fijo: van aparte como «pacc» (v1.14).
GF_KEYS = ["gfRRHH", "gfCargaFija", "gfAdmin", "gfMant", "gfITBIS", "gfProvision"]

GRUPO_CAMPO = {"LIC 2026-01": "lic2026", "LIC 2027-01": "lic2027", "ZFI 2027-02": "zfi", "Lotes adicionales": "adicionales",
               "CPAG HAB 2027": "cpag2027", "CPAG COM 2027": "cpag2027"}
ING_KEYS = ["cartera", "recupLegal", "lic2026", "lic2027", "zfi", "adicionales", "cpag2027"]
# Categoría de color en el mapa: los adjudicados de LIC 2026-01 ya están vendidos; los demás lotes de licitaciones
# dependen de la versión: «vende» si la versión los licita (y el escenario les da probabilidad), «noVende» si no.
GRUPO_CAT = {"LIC 2026-01": "adjudicada"}


def categoria_licitacion(l):
    return GRUPO_CAT.get(l["grupo"]) or ("vende" if l["incluida"] and l["probabilidad"] > 0 else "noVende")


def leer_panel(wb):
    """Indicadores del Panel. Desde v1.11 hay dos columnas: D = hasta dic-2027 y E = hasta el fin del
    horizonte (dic-2028). Devuelve (escenario, kpis del horizonte completo, kpis hasta dic-2027 o None)."""
    ws = wb["Panel"]
    escenario = ws["D5"].value
    labels = {norm(ws.cell(r, 2).value): r for r in range(1, 40) if ws.cell(r, 2).value}
    # fila de encabezado «Indicadores clave» (fila 7 hasta v1.18; fila 8 desde v1.19, que agregó el selector de versión)
    r_ind = labels.get(norm("Indicadores clave"), 7)
    dos_cols = norm(ws.cell(r_ind, 5).value).startswith("hasta")
    def leer(c):
        kpis = {}
        for key, alias in PANEL_KPIS:
            r = next((labels[norm(a)] for a in alias if norm(a) in labels), None)
            if r is None:
                fail(f"Panel: no encontré el indicador «{alias[0]}» en la columna B")
            v = ws.cell(r, c).value
            if key == "mesSaldoMin":
                kpis[key] = mes_key(v)
            elif key == "mesesNegativos":
                kpis[key] = int(num(v))
            else:
                kpis[key] = r2(num(v))
        return kpis
    if dos_cols:
        return escenario, leer(5), leer(4)
    return escenario, leer(4), None


def leer_flujo(wb):
    ws = wb["Flujo_Mensual"]
    if norm(ws["A4"].value) != "concepto":
        fail("Flujo_Mensual: se esperaba «Concepto» en A4")
    meses = meses_flujo(ws)
    filas = {}
    for r in range(5, ws.max_row + 1):
        lab = norm(ws.cell(r, 1).value)
        if lab:
            filas.setdefault(lab, r)
    out = {}
    for key, label in FLUJO_FILAS.items():
        alias = label if isinstance(label, tuple) else (label,)
        r = next((rr for a in alias for l, rr in filas.items() if l.startswith(norm(a))), None)
        if r is None and key in FLUJO_FILAS_OPCIONALES:
            out[key] = [0.0] * len(meses)
            continue
        if r is None:
            fail(f"Flujo_Mensual: no encontré la fila «{alias[0]}»")
        vals = [ws.cell(r, c).value for c in range(2, 2 + len(meses))]
        out[key] = [str(v) if v is not None else "" for v in vals] if key == "estado" else [num(v) for v in vals]
    return meses, out


def leer_obras(wb, meses):
    ws = wb["Obras"]
    post = etiqueta_posterior(ws)
    hr, hdr = find_header_row(ws, ["#", "Proyecto", "ID lote", "Clave mapa", "Obra", "Contratista", "Tipo", "Estatus",
                                   "Confianza del ID", "Contrato", "Adenda", "Restante por pagar", "Incluir",
                                   "Modo de programa", "Total horizonte", post])
    mc = month_cols(ws, hr, meses)
    S = "Obras"
    # v1.19: «Desplazamiento total» = Mover (meses) + retraso global de la versión + reprogramación de la versión
    c_desp = hdr.get(norm("Desplazamiento total"))
    out = []
    for r in data_rows(ws, hr + 1, col(hdr, "#", S)):
        g = lambda name: ws.cell(r, col(hdr, name, S)).value  # noqa: E731
        clave = str(g("Clave mapa") or "").strip()
        pagos = [num(ws.cell(r, c).value) for c in mc]
        out.append({
            "num": int(num(g("#"))),
            "desplazamiento": int(num(ws.cell(r, c_desp).value)) if c_desp else 0,
            "clave": clave,
            "proyecto": str(g("Proyecto") or "").strip(),
            "id": str(g("ID lote") or "").strip(),
            "obra": str(g("Obra") or "").strip(),
            "contratista": str(g("Contratista") or "").strip(),
            "tipo": str(g("Tipo") or "").strip(),
            "estatus": str(g("Estatus") or "").strip(),
            "confianza": str(g("Confianza del ID") or "").strip(),
            "contrato": r2(num(g("Contrato"))),
            "adenda": r2(num(g("Adenda"))),
            "restante": r2(num(g("Restante por pagar"))),
            "pagos": [r2(v) for v in pagos],
            "totalHorizonte": r2(num(g("Total horizonte"))),
            "posterior": r2(num(g(post))),
            "incluida": norm(g("Incluir")) in ("sí", "si"),
            "modo": str(g("Modo de programa") or "").strip(),
            "enMapa": bool(RE_CLAVE_REAL.match(clave)),
            "_pagosRaw": pagos,
        })
    if not out:
        fail("Obras: no hay filas de datos")
    return out


def leer_mantenimiento(wb, meses):
    """v1.14+: el total de la hoja «Mantenimiento» como una obra virtual de tipo Reserva (va en la fila 20 del flujo
    junto con la reserva de la hoja Obras). None en versiones anteriores."""
    if "Mantenimiento" not in wb.sheetnames:
        return None
    ws = wb["Mantenimiento"]
    mc = month_cols(ws, 4, meses)
    fila = next((r for r in range(1, ws.max_row + 1)
                 if norm(ws.cell(r, 2).value).startswith("total a cargo de vbc-rd")), None)
    if fila is None:
        fail("Mantenimiento: no encontré la fila «TOTAL A CARGO DE VBC-RD»")
    pagos = [num(ws.cell(fila, c).value) for c in mc]
    return {"num": 0, "desplazamiento": 0, "clave": "CJB-MANTENIMIENTO", "proyecto": "CJB", "id": "", "obra": "Mantenimiento de la ciudad (hoja Mantenimiento)",
            "contratista": "Varios", "tipo": "Reserva", "estatus": "", "confianza": "", "contrato": 0.0, "adenda": 0.0,
            "restante": 0.0, "pagos": [r2(v) for v in pagos], "totalHorizonte": r2(sum(pagos)), "posterior": 0.0,
            "incluida": True, "modo": "Hoja Mantenimiento", "enMapa": False, "_pagosRaw": pagos}


def leer_cuotas(wb, meses):
    ws = wb["Cuotas"]
    heads, rows = table_rows(ws, "tCuotas")
    # Desde v1.7 cada cuota se reparte por mes (retraso, renegociación, legal, terreno) en «Cobro aaaa-mm»
    cols_mes = [f"Cobro {m}" for m in meses]
    need = ["Grupo", "Proyecto", "IDLote", "ClaveMapa", "Cuota"] + cols_mes
    for h in need:
        if h not in heads:
            fail(f"Cuotas/tCuotas: falta la columna «{h}»")
    por_clave = defaultdict(lambda: zeros(len(meses)))
    for d in rows:
        for i, c in enumerate(cols_mes):
            por_clave[d["ClaveMapa"]][i] += num(d[c])
    return por_clave


def leer_licitaciones(wb, meses, cuotas, params=None):
    ws = wb["Licitaciones"]
    S = "Licitaciones"
    post = etiqueta_posterior(ws, "Cobro ")
    hr, hdr = find_header_row(ws, ["Grupo", "Proyecto", "ID lote", "Clave mapa", "Área m² (licitación)",
                                   "Precio aplicado RD$/m²", "Valor de venta", "Adjudicatario", "Fecha firma (lote)",
                                   "Incluir lote", "Probabilidad", "Valor esperado", "Cobro en horizonte", post])
    # v1.16+: la versión activa (Decisiones, sección 7) decide qué lotes se licitan y el área de los lotes de Santiago
    c_incl = hdr.get(norm("Se licita (versión activa)"), col(hdr, "Incluir lote", S))
    c_area = hdr.get(norm("Área aplicada m² (versión activa)"), col(hdr, "Área m² (licitación)", S))
    # v1.14+: tramo del calendario 3-4-3 de los lotes restantes de CJB (vacío en los demás)
    c_tramo = next((c for h, c in hdr.items() if h.startswith("tramo")), None)
    # Con «2 lotes» de Santiago (v1.16+), las filas de CPAG cuya área pasa a AreaLoteStgo son lotes genéricos «por
    # definir»: el modelo reutiliza las filas 01-01-01 y 02-01-01 sin que sean esos lotes, así que no se pintan en el mapa.
    params = params or {}
    stgo_gen = norm(params.get("lotesStgo")) == "2 lotes"
    area_stgo = num(params.get("areaLoteStgo"))
    n_gen = 0
    out = []
    for r in data_rows(ws, hr + 1, col(hdr, "Grupo", S)):
        g = lambda name: ws.cell(r, col(hdr, name, S)).value  # noqa: E731
        clave = str(g("Clave mapa") or "").strip()
        firma = g("Fecha firma (lote)")
        cobros = cuotas.get(clave, zeros(len(meses)))
        cobro_h = num(g("Cobro en horizonte"))
        if abs(sum(cobros) - cobro_h) > TOL:
            fail(f"Licitaciones {clave}: cuotas ({sum(cobros):,.2f}) ≠ cobro en horizonte ({cobro_h:,.2f})")
        incluida = norm(ws.cell(r, c_incl).value) in ("sí", "si")
        area = num(ws.cell(r, c_area).value)
        proyecto = str(g("Proyecto") or "").strip()
        lote = {
            "clave": clave,
            "proyecto": proyecto,
            "id": str(g("ID lote") or "").strip(),
            "grupo": str(g("Grupo") or "").strip(),
            "area": r2(area),
            "precio": r2(num(g("Precio aplicado RD$/m²"))),
            "valor": r2(num(g("Valor de venta"))),
            "adjudicatario": str(g("Adjudicatario") or "").strip(),
            "fechaFirma": firma.strftime("%Y-%m-%d") if isinstance(firma, (dt.datetime, dt.date)) else None,
            "incluida": incluida,
            "probabilidad": num(g("Probabilidad")),
            "valorEsperado": r2(num(g("Valor esperado"))),
            "cobros": [r2(v) for v in cobros],
            "cobroHorizonte": r2(cobro_h),
            "posterior": r2(num(g(post))),
            "tramo": int(num(ws.cell(r, c_tramo).value)) if c_tramo else 0,
            "enMapa": bool(RE_CLAVE_REAL.match(clave)),
            "_cobrosRaw": cobros,
        }
        if stgo_gen and proyecto == "CPAG" and incluida and area_stgo and abs(area - area_stgo) < 0.5:
            n_gen += 1
            lote.update(claveFuente=clave, clave=f"CPAG-STGO-{n_gen}", id=f"Lote {n_gen} por definir",
                        generico=True, enMapa=False)
        out.append(lote)
    return out


def leer_mes_recup_legal(wb):
    ws = wb["Decisiones"]
    for r in range(1, ws.max_row + 1):
        if norm(ws.cell(r, 2).value).startswith("mes estimado de esa recuperación"):
            return mes_key(ws.cell(r, 3).value)
    fail("Decisiones: no encontré «Mes estimado de esa recuperación»")


def leer_cartera(wb, meses):
    ws = wb["CxC_Cartera"]
    S = "CxC_Cartera"
    post = etiqueta_posterior(ws)
    hr, hdr = find_header_row(ws, ["Proyecto", "ID lote", "Clave mapa", "Cliente (archivo)", "Nombre (mapa)",
                                   "Saldo según archivo", "Situación", "Recuperación legal esperada",
                                   "Total horizonte", post])
    mc = month_cols(ws, hr, meses)
    mes_legal = leer_mes_recup_legal(wb)
    out = []
    for r in data_rows(ws, hr + 1, col(hdr, "Proyecto", S)):
        g = lambda name: ws.cell(r, col(hdr, name, S)).value  # noqa: E731
        clave = str(g("Clave mapa") or "").strip()
        cobros = [num(ws.cell(r, c).value) for c in mc]
        recup = num(g("Recuperación legal esperada"))
        out.append({
            "clave": clave,
            "proyecto": str(g("Proyecto") or "").strip(),
            "id": str(g("ID lote") or "").strip(),
            "cliente": str(g("Cliente (archivo)") or "").strip(),
            "nombre": str(g("Nombre (mapa)") or "").strip(),
            "situacion": str(g("Situación") or "").strip(),
            "saldo": r2(num(g("Saldo según archivo"))),
            "cobros": [r2(v) for v in cobros],
            "totalHorizonte": r2(num(g("Total horizonte"))),
            "posterior": r2(num(g(post))),
            "recupLegal": r2(recup),
            "mesRecupLegal": mes_legal if recup else None,
            "enMapa": bool(RE_CLAVE_REAL.match(clave)),
            "_cobrosRaw": cobros,
            "_recupRaw": recup,
        })
    return out, mes_legal


def leer_retenidos(wb, meses):
    ws = wb["Retenidos"]
    S = "Retenidos"
    post = etiqueta_posterior(ws)
    hr, hdr = find_header_row(ws, ["#", "Proyecto", "ID lote", "Clave mapa", "Contratista", "Obra", "Monto retenido",
                                   "Mes de pago (calculado)", "Total horizonte", post])
    mc = month_cols(ws, hr, meses)
    out = []
    for r in data_rows(ws, hr + 1, col(hdr, "#", S)):
        g = lambda name: ws.cell(r, col(hdr, name, S)).value  # noqa: E731
        clave = str(g("Clave mapa") or "").strip()
        pagos = [num(ws.cell(r, c).value) for c in mc]
        out.append({
            "clave": clave,
            "proyecto": str(g("Proyecto") or "").strip(),
            "id": str(g("ID lote") or "").strip(),
            "contratista": str(g("Contratista") or "").strip(),
            "obra": str(g("Obra") or "").strip(),
            "monto": r2(num(g("Monto retenido"))),
            "mes": mes_key(g("Mes de pago (calculado)")),
            "pagos": [r2(v) for v in pagos],
            "totalHorizonte": r2(num(g("Total horizonte"))),
            "posterior": r2(num(g(post))),
            "enMapa": bool(RE_CLAVE_REAL.match(clave)),
            "_pagosRaw": pagos,
        })
    return out


def leer_maestro(wb):
    ws = wb["Lotes_Maestro"]
    heads, rows = table_rows(ws, "tLotes")
    need = ["Proyecto", "IDLote", "ClaveMapa", "Tipo", "Subtipo", "Nombre", "Desarrollador", "AreaM2", "Estado",
            "IngresosHorizonte", "EgresosObrasHorizonte", "FlujoNetoLote", "ConMovimiento"]
    for h in need:
        if h not in heads:
            fail(f"Lotes_Maestro/tLotes: falta la columna «{h}»")
    return rows


# ----------------------------------------------------------------------------- versiones (v1.19+)
VERSION_COLS = [  # clave JSON, inicio del encabezado en Decisiones, sección 8
    ("antes", "Antes"), ("nLotes", "Lotes que se licitan"), ("lotesCJB", "Lotes restantes de CJB"),
    ("lotesStgo", "Lotes de Santiago"), ("precio", "Precio por m²"), ("calendario", "Calendario"),
    ("pctVendido", "% que se vende"), ("retrasoObras", "Retraso global de obras"), ("reprogramadas", "Obras reprogramadas"),
]
VERSION_NUM = {"nLotes", "precio", "pctVendido", "retrasoObras"}


def leer_versiones(wb):
    """Versiones de ventas y obras (Decisiones, sección 8) en el orden de la lista del Panel D7, y la tabla de obras que
    mueve la «reprogramación». ([], []) en modelos sin selector de versión (antes de v1.19)."""
    if "Decisiones" not in wb.sheetnames or celda_nombre(wb, "VersionActiva") is None:
        return [], []
    ws = wb["Decisiones"]
    S = "Decisiones"
    filas = filas_etiqueta(ws)
    hr = buscar_fila(filas, "Versión (lista del Panel)", S)
    cols = columnas_encabezado(ws, hr, VERSION_COLS, S, c0=3)
    versiones = []
    for r in range(hr + 1, hr + 20):
        et = ws.cell(r, 2).value
        if not et or norm(et).startswith("personalizada"):
            break
        m = re.match(r"^\s*(\d+)\s*·\s*(.+?)\s*$", str(et))
        if not m:
            fail(f"Decisiones B{r}: la versión «{et}» no empieza por «n · »")
        d = {"n": int(m.group(1)), "etiqueta": str(et).strip(), "nombre": m.group(2)}
        for k, c in cols.items():
            v = ws.cell(r, c).value
            d[k] = num(v) if k in VERSION_NUM else str(v or "").strip()
        d["reprogramadas"] = norm(d["reprogramadas"]) in ("sí", "si")
        versiones.append(d)
    reprog = []
    rr = buscar_fila(filas, "Obra (hoja Obras)", S, opcional=True)
    if rr:
        for r in range(rr + 1, rr + 20):
            if not ws.cell(r, 2).value:
                break
            reprog.append({"obra": str(ws.cell(r, 2).value).strip(), "num": int(num(ws.cell(r, 3).value)),
                           "meses": num(ws.cell(r, 4).value)})
    return versiones, reprog


# Hoja «Versiones»: textos que describen cada versión y valores fijos calculados con la Base (se comparan al exportar)
VERSIONES_DESC = [("precioTxt", "Ventas: precio por m²"), ("lotesTxt", "Ventas: qué lotes se licitan y cuándo"),
                  ("vendidoTxt", "Ventas: % de los lotes que se venden"), ("obrasTxt", "Pagos de obra")]
VERSIONES_FIJOS = [("ingresos2027", "Ingresos totales 2027"), ("egresos2027", "Egresos totales 2027"),
                   ("saldoMin2027", "Saldo de caja más bajo de 2027"), ("saldoDic27", "Saldo de caja a dic-2027"),
                   ("saldoDic28", "Saldo de caja a dic-2028"), ("capacidad", "… hasta dic-27 (capacidad 2027)"),
                   ("techo", "Techo global de compras nuevas (el menor"), ("santiago", "Reserva para Santiago"),
                   ("ratio", "No-obra / obra con compras y Santiago")]


def leer_hoja_versiones(wb):
    """{n: {"desc": {...}, "fijos": {...}}} de la hoja «Versiones» (v1.19); {} si no existe."""
    if "Versiones" not in wb.sheetnames:
        return {}
    ws = wb["Versiones"]
    filas = filas_etiqueta(ws)
    hr = buscar_fila(filas, "Concepto", "Versiones", exacta=True)
    cols = {}
    for c in range(3, ws.max_column + 1):
        m = re.match(r"^versión (\d+)$", norm(ws.cell(hr, c).value))
        if m:
            cols[int(m.group(1))] = c
    out = {n: {"desc": {}, "fijos": {}} for n in cols}
    for grupo, lista in (("desc", VERSIONES_DESC), ("fijos", VERSIONES_FIJOS)):
        for clave, ini in lista:
            r = buscar_fila(filas, ini, "Versiones", opcional=True)
            if r is None:
                continue
            for n, c in cols.items():
                v = ws.cell(r, c).value
                out[n][grupo][clave] = v if grupo == "fijos" else valor_celda(v, ratio=clave == "vendidoTxt")
    return out


# Supuestos activos de cada combinación, leídos por nombre definido (los que no existan en el modelo se omiten)
PARAMS = [
    ("versionN", "VersionN"), ("escenario", "Escenario"), ("fuenteCub", "FuenteCub"), ("ret1Remitida", "Ret1Remitida"),
    ("retrasoObras", "RetrasoCub"), ("precioRestantes", "PrecioRestantes"), ("calendario", "CalendarioRestantes"),
    ("mesesEntreTramos", "MesesEntreTramos"), ("probLic2027", "ProbLic2027"), ("lotesCJB", "LotesCJB"),
    ("lotesStgo", "LotesStgo"), ("areaLoteStgo", "AreaLoteStgo"), ("precioLoteStgo", "PrecioLoteStgo"),
    ("ajusteGastos", "AjusteGastos"), ("colchon", "Colchon"), ("saldoInicial", "SaldoInicial"), ("topeNoObra", "TopeNoObra"),
]


def leer_parametros(wb):
    out = {}
    for k, nombre in PARAMS:
        c = celda_nombre(wb, nombre)
        if c is not None:
            out[k] = valor_celda(c.value, ratio=k in ("probLic2027", "ajusteGastos", "topeNoObra"))
    return out


# ----------------------------------------------------------------------------- 75/25, techos y compras (v1.15+)
SPLIT_VALORES = [  # clave JSON, inicio de la etiqueta en Split_Santiago (columna B; el valor está en C)
    ("tope", "Tope del gasto no-obra"), ("mantEsObra", "¿El mantenimiento de la ciudad cuenta como obra?"),
    ("reservaMin", "Reserva mínima para Santiago"), ("reservaTope", "Tope de la reserva para Santiago"),
    ("ratioAntes", "Gasto no-obra / obra antes de compras y Santiago"),
    ("disponible", "Dinero disponible a dic-2027 sin bajar del colchón"),
    ("techoReglaTodoStgo", "Techo que permite la regla si todo lo demás va a Santiago"),
    ("techoReglaStgoTope", "Techo que permite la regla si Santiago llega a su tope"),
    ("techoReservaMin", "Techo que deja libre la reserva mínima de Santiago"),
    ("techoRegla", "Techo según la regla del split"), ("techoPolitica", "Techo de política"),
    ("techoGlobal", "TECHO GLOBAL DE COMPRAS DE LAS ÁREAS"), ("reservaSantiago", "RESERVA PARA SANTIAGO"),
    ("sinAsignar", "Sin asignar"), ("noObraConCompras", "Gasto no-obra 2027 con las compras"),
    ("obraConSantiago", "Obra 2027 con Santiago"), ("ratio", "Gasto no-obra / obra (resultado)"),
    ("cumple", "¿Cumple el tope?"), ("recorte", "Recorte de gasto no-obra que haría falta para cumplir con compras en 0"),
    ("saldoDespues", "Saldo de caja a dic-2027 después de compras y Santiago"),
    ("recorteTodoStgo", "Recorte de gasto no-obra para cumplir con compras en 0 y todo el disponible en Santiago"),
    ("necesidadAreas", "Necesidad recurrente de compras de las áreas"),
    ("recorteNecesidad", "Recorte de gasto no-obra para darle a las áreas esa necesidad completa"),
    ("techoEstricto", "Techo de compras con el criterio estricto"), ("stgoEstricto", "… y reserva para Santiago en ese caso"),
]
SPLIT_RATIOS = {"tope", "ratioAntes", "ratio"}
SPLIT_OPCIONALES = {"reservaMin", "reservaTope", "techoReglaTodoStgo", "techoReglaStgoTope", "techoReservaMin",
                    "recorteTodoStgo", "necesidadAreas", "recorteNecesidad", "techoEstricto", "stgoEstricto"}
SPLIT_MES = [("capacidadAcum", "Capacidad acumulada sin bajar del colchón"),
             ("comprasAcum", "Compras de las áreas liberadas, acumulado"), ("holgura", "Holgura después de las compras"),
             ("stgoAcum", "Santiago: se puede comprometer, acumulado"), ("stgoMes", "Santiago: nuevo en el mes")]


def leer_split(wb):
    """Hoja Split_Santiago: regla 75/25, reparto del disponible entre compras y Santiago, qué falta para cumplir,
    cuándo hay dinero y en qué se usaría la reserva de Santiago. None en modelos anteriores a v1.15."""
    if "Split_Santiago" not in wb.sheetnames:
        return None
    ws = wb["Split_Santiago"]
    S = "Split_Santiago"
    filas = filas_etiqueta(ws)
    out = {}
    for k, ini in SPLIT_VALORES:
        r = buscar_fila(filas, ini, S, opcional=k in SPLIT_OPCIONALES)
        if r is not None:
            out[k] = valor_celda(ws.cell(r, 3).value, ratio=k in SPLIT_RATIOS)
    # sección 2: cada línea del flujo 2027 clasificada como obra o no-obra
    hr = buscar_fila(filas, "Línea de Flujo_Mensual", S)
    cc = columnas_encabezado(ws, hr, [("clase", "Clase"), ("monto", "Año 2027")], S)
    out["lineas"] = []
    for r in range(hr + 1, hr + 40):
        lab = ws.cell(r, 2).value
        if not lab or norm(lab).startswith("total"):
            break
        out["lineas"].append({"linea": str(lab).strip(), "clase": str(ws.cell(r, cc["clase"]).value or "").strip(),
                              "monto": r2(num(ws.cell(r, cc["monto"]).value))})
    # sección 5: cuándo hay dinero en 2027 (acumulado)
    rm = buscar_fila(filas, "Mes", S, opcional=True, exacta=True)
    if rm:
        cm = columnas_mes(ws, rm)
        out["meses"] = [m for m, _ in cm]
        sub = [f for f in filas if f[1] > rm]
        for k, ini in SPLIT_MES:
            r = buscar_fila(sub, ini, S, opcional=True)
            if r:
                out[k] = [r2(num(ws.cell(r, c).value)) for _, c in cm]
    # sección 6: uso propuesto de la reserva de Santiago
    r6 = buscar_fila(filas, "6. Uso propuesto de la reserva", S, opcional=True)
    if r6:
        hr6 = buscar_fila([f for f in filas if f[1] > r6], "Concepto", S, exacta=True)
        c6 = columnas_encabezado(ws, hr6, [("monto", "Monto estimado"), ("asignado", "Asignado con la reserva"),
                                           ("cuando", "Cuándo"), ("condicion", "Condición")], S)
        out["usosSantiago"] = []
        for r in range(hr6 + 1, hr6 + 30):
            lab = ws.cell(r, 2).value
            if not lab or norm(lab) == "total":
                break
            out["usosSantiago"].append({"concepto": str(lab).strip(), "monto": r2(num(ws.cell(r, c6["monto"]).value)),
                                        "asignado": r2(num(ws.cell(r, c6["asignado"]).value)),
                                        "cuando": str(ws.cell(r, c6["cuando"]).value or "").strip(),
                                        "condicion": str(ws.cell(r, c6["condicion"]).value or "").strip()})
    return out


TECHOS_VALORES = [("capacidad", "Capacidad de caja para compras 2027"), ("politica", "Techo de política"),
                  ("regla", "Techo que permite la regla 75/25"), ("global", "Techo global (el menor de los tres)"),
                  ("aRepartir", "A repartir entre las áreas"), ("pisos", "Suma de los pisos de las áreas"),
                  ("factorPisos", "Factor de los pisos"), ("asignado", "Techo asignado"),
                  ("ejecutable", "Ejecutable por caja en 2027"), ("noAsignado", "No asignado")]
TECHOS_MES = [("saldoAntes", "Saldo antes de las compras"), ("colchon", "Colchón mínimo"),
              ("capacidadAcum", "Capacidad acumulada para compras"), ("liberacionPlan", "Liberación planificada acumulada"),
              ("liberacionPermitida", "Liberación permitida acumulada"),
              ("comprasMes", "Compras de las áreas liberadas en el mes"),
              ("saldoDespues", "Saldo final de caja después de las compras")]
AREA_COLS = [("tipo", "Tipo"), ("empleados", "Empleados actuales"), ("plazas", "Plazas nuevas 2027"),
             ("empleados2027", "Empleados 2027"), ("indice", "Índice de operaciones"),
             ("justificacion", "Justificación del índice"), ("pacc2025", "PACC 2025"), ("pacc2026", "PACC 2026"),
             ("base", "Base histórica ajustada"), ("piso", "Piso"), ("techo", "Techo anual 2027"),
             ("t1", "T1 ("), ("t2", "T2 ("), ("t3", "T3 ("), ("t4", "T4 ("), ("tecnologia", "De lo cual tecnología"),
             ("pacOperacion", "PAC de operación que ya pagan"), ("pacTotal", "PAC total del área 2027")]
AREA_TEXTO = {"tipo", "justificacion"}


def leer_techos(wb):
    """Hoja Techos_Areas: capacidad de caja para compras mes a mes (sección 2), techo global (sección 3) y techo por
    área (sección 4). None en modelos anteriores a v1.14."""
    if "Techos_Areas" not in wb.sheetnames:
        return None
    ws = wb["Techos_Areas"]
    S = "Techos_Areas"
    filas = filas_etiqueta(ws)
    r3 = buscar_fila(filas, "3. Techo global", S)
    sub3 = [f for f in filas if f[1] > r3]
    out = {k: valor_celda(ws.cell(buscar_fila(sub3, ini, S), 3).value, ratio=k == "factorPisos")
           for k, ini in TECHOS_VALORES}
    r2_ = buscar_fila(filas, "2. Capacidad de caja", S)
    hr = buscar_fila([f for f in filas if r2_ < f[1] < r3], "Concepto", S, exacta=True)
    cm = columnas_mes(ws, hr)
    out["meses"] = [m for m, _ in cm]
    sub2 = [f for f in filas if hr < f[1] < r3]
    for k, ini in TECHOS_MES:
        r = buscar_fila(sub2, ini, S)
        out[k] = [r2(num(ws.cell(r, c).value)) for _, c in cm]
    hr4 = buscar_fila(filas, "Dirección / área", S)
    cc = columnas_encabezado(ws, hr4, AREA_COLS, S, c0=3)
    out["areas"] = []
    for r in range(hr4 + 1, hr4 + 40):
        nom = ws.cell(r, 2).value
        if not nom or norm(nom).startswith("total"):
            break
        a = {"area": str(nom).strip()}
        for k, c in cc.items():
            v = ws.cell(r, c).value
            a[k] = str(v or "").strip() if k in AREA_TEXTO else r2(num(v))
        out["areas"].append(a)
    return out


PANEL_EXTRA = [("techoCompras", "Techo global de compras nuevas de las áreas 2027"),
               ("capacidadCompras", "Capacidad de caja para compras 2027"),
               ("reservaSantiago", "Reserva para Santiago (CPAG) 2027"),
               ("ratioNoObra", "Gasto no-obra / obra 2027 (con compras y Santiago)")]


def leer_panel_extra(wb):
    """Panel v1.15+: compras, Santiago y 75/25 (columna D), la tabla «¿Alcanza para repartir…?» (v1.19: D = techo del
    modelo, E = monto a evaluar) y el impacto de riesgos y palancas del escenario."""
    ws = wb["Panel"]
    filas = filas_etiqueta(ws, 2, 1, 120)
    out = {}
    for k, ini in PANEL_EXTRA:
        r = buscar_fila(filas, ini, "Panel", opcional=True)
        if r:
            out[k] = valor_celda(ws.cell(r, 4).value, ratio=k == "ratioNoObra")
    r = buscar_fila(filas, "¿Alcanza para repartir", "Panel", opcional=True)
    if r:
        hr = buscar_fila([f for f in filas if f[1] > r], "Concepto", "Panel", exacta=True)
        out["alcanza"] = []
        for rr in range(hr + 1, hr + 25):
            lab = ws.cell(rr, 2).value
            if not lab or norm(lab).startswith("palanca: vender"):
                break
            ratio = norm(lab).startswith("gasto no-obra")
            out["alcanza"].append({"concepto": str(lab).strip(), "modelo": valor_celda(ws.cell(rr, 4).value, ratio),
                                   "evaluar": valor_celda(ws.cell(rr, 5).value, ratio),
                                   "nota": str(ws.cell(rr, 6).value or "").strip()})
    r = buscar_fila(filas, "Impacto de riesgos y palancas", "Panel", opcional=True)
    if r:
        out["impactos"] = []
        for rr in range(r + 1, r + 15):
            lab = ws.cell(rr, 2).value
            if not lab:
                break
            out["impactos"].append({"concepto": str(lab).strip(), "monto": r2(num(ws.cell(rr, 4).value)),
                                    "nota": str(ws.cell(rr, 5).value or "").strip()})
    return out


COMPRAS_COMBO = [  # Compras_Historicas, secciones 4 y 5 (dependen del escenario: ajuste de gastos, y de la versión)
    ("criterio", "Criterio"), ("pacRef", "PAC (o bienes, con el 3.er criterio) del año de referencia"),
    ("inflacion", "Inflación anual de precios de compra"), ("anios", "Años hasta 2027"),
    ("pacRef2027", "Monto del año de referencia a precios de 2027"), ("plazas", "Plazas nuevas 2027"),
    ("equipoPlaza", "Equipo por plaza nueva"), ("equipoPlazas", "Equipo de las plazas nuevas"), ("redondeo", "Redondeo"),
    ("historial2027", "Historial a precios de 2027 + equipo"),
    ("pacEnGastosFijos", "Menos: PAC de operación que ya pagan los gastos fijos"),
    ("pacTotal", "PAC TOTAL 2027 POR HISTORIAL"), ("techoHistorial", "TECHO DE COMPRAS NUEVAS POR HISTORIAL"),
    ("techoV118", "Antes: v1.18"), ("necesidadAreas", "Necesidad limpia de las áreas"),
    ("techoRegla", "Techo que permite la regla 75/25"), ("capacidad", "Capacidad de caja para compras 2027"),
    ("techoGlobal", "TECHO GLOBAL DE COMPRAS NUEVAS 2027"), ("pacTotalModelo", "PAC total 2027 (operación que ya pagan"),
    ("reservaSantiago", "Reserva para Santiago (lo que queda"), ("ratio", "No-obra / obra con compras y Santiago"),
    ("ratioSinStgo", "No-obra / obra sin Santiago (si Santiago no pagara"),
    ("obraMinStgo", "Obra mínima en Santiago en 2027"),
    ("ratioRitmo2026", "No-obra / obra sin Santiago con los gastos fijos al ritmo real"),
]
COMPRAS_RATIOS = {"inflacion", "ratio", "ratioSinStgo", "ratioRitmo2026"}


def leer_compras_combo(wb):
    """Compras_Historicas, secciones 4 (techo por historial) y 5 (cómo queda 2027). None antes de v1.18."""
    if "Compras_Historicas" not in wb.sheetnames:
        return None
    ws = wb["Compras_Historicas"]
    S = "Compras_Historicas"
    filas = filas_etiqueta(ws)
    r4 = buscar_fila(filas, "4. Techo de compras 2027", S)
    sub = [f for f in filas if f[1] > r4]
    out = {}
    for k, ini in COMPRAS_COMBO:
        r = buscar_fila(sub, ini, S, opcional=True)
        if r:
            out[k] = valor_celda(ws.cell(r, 3).value, ratio=k in COMPRAS_RATIOS)
    return out


def leer_compras_hist(wb):
    """Compras_Historicas, partes que no dependen de la versión ni del escenario: bienes del registro de activos por año
    (sección 1), pagos a suplidores por línea con su clase (sección 2) y PAC real por año (sección 3)."""
    if "Compras_Historicas" not in wb.sheetnames:
        return None
    ws = wb["Compras_Historicas"]
    S = "Compras_Historicas"
    filas = filas_etiqueta(ws)
    out = {}
    # 1. bienes por año (registro de activos fijos)
    hr = buscar_fila(filas, "Clase del registro", S)
    anios = [(norm(ws.cell(hr, c).value)[:4], c) for c in range(3, ws.max_column + 1)
             if re.fullmatch(r"20\d\d(\.0)?", norm(ws.cell(hr, c).value))]
    c_que = columnas_encabezado(ws, hr, [("que", "Qué se compró")], S)["que"]
    out["bienesAnios"] = [a for a, _ in anios]
    out["bienes"] = []
    for r in range(hr + 1, hr + 30):
        lab = ws.cell(r, 2).value
        if not lab or norm(lab).startswith("total"):
            break
        out["bienes"].append({"clase": str(lab).strip(), "valores": [r2(num(ws.cell(r, c).value)) for _, c in anios],
                              "detalle": str(ws.cell(r, c_que).value or "").strip()})
    # 2. pagos del flujo de Finanzas por suplidor o línea
    hr = buscar_fila(filas, "Suplidor o línea", S)
    cc = columnas_encabezado(ws, hr, [("que", "Qué es"), ("clase", "Clase"), ("area", "Área que lo compra"),
                                      ("y24", "2024"), ("y25", "2025"), ("y26a", "2026 ene"), ("y26b", "2026 oct"),
                                      ("nota", "Nota")], S)
    out["periodos"] = [str(ws.cell(hr, cc[k]).value).strip() for k in ("y24", "y25", "y26a", "y26b")]
    out["lineas"] = []
    for r in range(hr + 1, hr + 80):
        lab = ws.cell(r, 2).value
        if not lab or norm(lab).startswith("total de las líneas"):
            break
        out["lineas"].append({"linea": str(lab).strip(), "que": str(ws.cell(r, cc["que"]).value or "").strip(),
                              "clase": str(ws.cell(r, cc["clase"]).value or "").strip(),
                              "area": str(ws.cell(r, cc["area"]).value or "").strip(),
                              "valores": [r2(num(ws.cell(r, cc[k]).value)) for k in ("y24", "y25", "y26a", "y26b")],
                              "nota": str(ws.cell(r, cc["nota"]).value or "").strip()})
    sub = [f for f in filas if f[1] > hr]
    out["resumen"] = {}
    for k, ini in (("bienes", "PAC: bienes y tecnología"), ("operacion", "PAC: operación y servicios"),
                   ("total", "PAC total (las dos filas anteriores)"), ("fueraPAC", "Fuera del PAC")):
        r = buscar_fila(sub, ini, S)
        out["resumen"][k] = [r2(num(ws.cell(r, cc[c]).value)) for c in ("y24", "y25", "y26a", "y26b")]
    # 3. PAC real por año
    r3 = buscar_fila(filas, "3. PAC real por año", S)
    hr = buscar_fila([f for f in filas if f[1] > r3], "Concepto", S, exacta=True)
    cols = [(norm(ws.cell(hr, c).value).removesuffix(".0"), c) for c in range(3, ws.max_column + 1) if ws.cell(hr, c).value]
    out["porAnio"] = {"periodos": [p for p, _ in cols], "filas": []}
    for r in range(hr + 1, hr + 15):
        lab = ws.cell(r, 2).value
        if not lab:
            break
        ratio = norm(lab).startswith("pagado / plan")
        valores = [None if isinstance(ws.cell(r, c).value, str) or ws.cell(r, c).value is None   # «n/d» → null
                   else valor_celda(ws.cell(r, c).value, ratio) for _, c in cols]
        if any(v is not None for v in valores):   # las notas al pie de la tabla no tienen cifras
            out["porAnio"]["filas"].append({"concepto": str(lab).strip(), "valores": valores})
    return out


def leer_gasto_real(wb):
    """Gasto_Real_2026 (v1.17+), lo que no depende de la versión: presupuesto 2026 frente a lo pagado por categoría,
    obra / no-obra pagados por año, partidas del no-obra y mes a mes de 2026."""
    if "Gasto_Real_2026" not in wb.sheetnames:
        return None
    ws = wb["Gasto_Real_2026"]
    S = "Gasto_Real_2026"
    filas = filas_etiqueta(ws)
    out = {}
    hr = buscar_fila(filas, "Categoría", S, exacta=True)
    cc = columnas_encabezado(ws, hr, [("presAnio", "Presupuesto 2026 (año)"), ("presEneSep", "Presupuesto ene"),
                                      ("pagado", "Pagado ene"), ("ritmo", "Al ritmo de ene")], S)
    out["presupuesto"] = []
    for r in range(hr + 1, hr + 20):
        lab = ws.cell(r, 2).value
        if not lab or norm(lab).startswith("el presupuesto"):
            break
        out["presupuesto"].append({"categoria": str(lab).strip(), **{k: r2(num(ws.cell(r, c).value)) for k, c in cc.items()}})
    hr = buscar_fila(filas, "Período", S, exacta=True)
    cc = columnas_encabezado(ws, hr, [("obra", "Obra"), ("noObra", "No-obra"), ("ratio", "No-obra / obra")], S)
    out["porAnio"] = []
    for r in range(hr + 1, hr + 15):
        lab = ws.cell(r, 2).value
        if not lab or not re.match(r"^20\d\d", str(lab)):
            break
        if "modelo" in norm(lab):
            continue  # 2027 del modelo: sale de cada combinación (Split_Santiago)
        out["porAnio"].append({"periodo": str(lab).strip(), "obra": r2(num(ws.cell(r, cc["obra"]).value)),
                               "noObra": r2(num(ws.cell(r, cc["noObra"]).value)),
                               "ratio": round(num(ws.cell(r, cc["ratio"]).value), 6)})
    hr = buscar_fila(filas, "Partida", S, exacta=True)
    cc = columnas_encabezado(ws, hr, [("y25", "2025 pagado"), ("y26", "2026 ene"), ("y26b", "2026 oct"),
                                      ("y26t", "2026 año")], S)
    out["partidas"] = []
    for r in range(hr + 1, hr + 30):
        lab = ws.cell(r, 2).value
        if not lab or norm(lab).startswith("memo") or norm(lab).startswith("control"):
            break
        out["partidas"].append({"partida": str(lab).strip(), **{k: r2(num(ws.cell(r, c).value)) for k, c in cc.items()}})
    r5 = buscar_fila(filas, "5. Mes a mes en 2026", S, opcional=True)
    if r5:
        sub = [f for f in filas if f[1] > r5]
        hr = buscar_fila(sub, "Concepto", S, exacta=True)
        cm = columnas_mes(ws, hr)
        out["mensual"] = {"meses": [m for m, _ in cm]}
        for k, ini in (("tipo", "Tipo"), ("obra", "Obra"), ("noObra", "No-obra"), ("cobros", "Cobro de lotes")):
            r = buscar_fila(sub, ini, S, exacta=True)
            out["mensual"][k] = [str(ws.cell(r, c).value or "") if k == "tipo" else r2(num(ws.cell(r, c).value)) for _, c in cm]
    return out


def leer_riesgos(wb):
    """Panel, «Supuestos de cobro y riesgos»: parámetros de cada escenario (Pesimista, Base, Optimista)."""
    ws = wb["Panel"]
    filas = filas_etiqueta(ws, 2, 1, 120)
    r = buscar_fila(filas, "Supuestos de cobro y riesgos", "Panel", opcional=True)
    if not r:
        return []
    hr = buscar_fila([f for f in filas if f[1] > r], "Parámetro", "Panel", exacta=True)
    cc = columnas_encabezado(ws, hr, [(e, e) for e in ESCENARIOS] + [("explicacion", "Explicación")], "Panel")
    out = []
    for rr in range(hr + 1, hr + 20):
        lab = ws.cell(rr, 2).value
        if not lab or norm(lab).startswith("lote con caso"):
            break
        d = {"parametro": str(lab).strip(), "explicacion": str(ws.cell(rr, cc["explicacion"]).value or "").strip()}
        for e in ESCENARIOS:
            d[e] = round(num(ws.cell(rr, cc[e]).value), 6)
        out.append(d)
    return out


# ----------------------------------------------------------------------------- armado
CATEGORIAS = ["obra", "adjudicada", "vende", "noVende", "cartera", "legal"]


def armar_escenario(wb, meses):
    nombre, kpis, kpis2027 = leer_panel(wb)
    meses_f, f = leer_flujo(wb)
    if meses_f != meses:
        fail("Los meses de Flujo_Mensual no coinciden entre escenarios")
    params = leer_parametros(wb)
    obras = leer_obras(wb, meses)
    mant = leer_mantenimiento(wb, meses)
    if mant:
        obras.append(mant)
    cuotas = leer_cuotas(wb, meses)
    lics = leer_licitaciones(wb, meses, cuotas, params)
    cartera, _ = leer_cartera(wb, meses)
    retenidos = leer_retenidos(wb, meses)
    maestro = leer_maestro(wb)
    n = len(meses)
    idx = {m: i for i, m in enumerate(meses)}

    # --- cuadre de detalles contra Flujo_Mensual
    def check(nombre_fila, calc, ref):
        for i, (a, b) in enumerate(zip(calc, ref)):
            if abs(a - b) > TOL:
                fail(f"No cuadra «{nombre_fila}» en {meses[i]}: detalle {a:,.2f} vs Flujo_Mensual {b:,.2f}")

    acc = defaultdict(lambda: zeros(n))
    for o in obras:
        if o["tipo"].lower() == "reserva":
            add_into(acc["mantGeneral"], o["_pagosRaw"])
        else:
            add_into(acc["obras" + o["proyecto"]], o["_pagosRaw"])
    for l in lics:
        add_into(acc[GRUPO_CAMPO[l["grupo"]]], l["_cobrosRaw"])
    for c in cartera:
        add_into(acc["cartera"], c["_cobrosRaw"])
        if c["_recupRaw"]:
            acc["recupLegal"][idx[c["mesRecupLegal"]]] += c["_recupRaw"]
    for rt in retenidos:
        add_into(acc["retenidos"], rt["_pagosRaw"])
    for k in ["obrasCJB", "obrasCPAG", "mantGeneral", "retenidos"] + ING_KEYS:
        check(k, acc[k], f[k])

    gastos_fijos = [sum(v) for v in zip(*(f[k] for k in GF_KEYS))]
    tot_ing = [sum(v) for v in zip(*(f[k] for k in ING_KEYS + ["inversiones", "otros"]))]
    tot_egr = [sum(v) for v in zip(f["obrasCJB"], f["obrasCPAG"], f["mantGeneral"], f["complementarias"],
                                   f["costoAtraso"], f["retenidos"], gastos_fijos, f["obligaciones"], f["gfBienes"])]
    check("TOTAL INGRESOS", tot_ing, f["totalIngresos"])
    check("TOTAL EGRESOS", tot_egr, f["totalEgresos"])

    # --- cuadre contra el Panel
    for key, serie in (("ingresos", f["totalIngresos"]), ("egresos", f["totalEgresos"])):
        if abs(sum(serie) - kpis[key]) > TOL:
            fail(f"Panel {key} ({kpis[key]:,.2f}) ≠ suma de Flujo_Mensual ({sum(serie):,.2f})")
    if abs(f["saldo"][-1] - kpis["saldoFinal"]) > TOL:
        fail("Panel saldo final ≠ Flujo_Mensual")
    if abs(min(f["saldo"]) - kpis["saldoMin"]) > TOL:
        fail("Panel saldo mínimo ≠ Flujo_Mensual")
    if kpis2027:   # columna «Hasta dic-2027» del Panel (v1.11+)
        n27 = meses.index("2027-12") + 1
        for key, serie in (("ingresos", f["totalIngresos"]), ("egresos", f["totalEgresos"])):
            if abs(sum(serie[:n27]) - kpis2027[key]) > TOL:
                fail(f"Panel {key} hasta dic-2027 ({kpis2027[key]:,.2f}) ≠ Flujo_Mensual ({sum(serie[:n27]):,.2f})")
        if abs(f["saldo"][n27 - 1] - kpis2027["saldoFinal"]) > TOL or abs(min(f["saldo"][:n27]) - kpis2027["saldoMin"]) > TOL:
            fail("Panel saldo hasta dic-2027 ≠ Flujo_Mensual")

    # --- categorías y totales por lote (solo claves reales)
    lotes = defaultdict(lambda: {"categorias": [], "ingresos": 0.0, "egresos": 0.0,
                                 "ingresosMes": zeros(n), "egresosMes": zeros(n), "tramo": 0, "desplazamiento": 0})

    def tag(clave, cat):
        if RE_CLAVE_REAL.match(clave) and cat not in lotes[clave]["categorias"]:
            lotes[clave]["categorias"].append(cat)

    for o in obras:
        if o["enMapa"]:
            tag(o["clave"], "obra")
            add_into(lotes[o["clave"]]["egresosMes"], o["_pagosRaw"])
            # meses que la versión mueve los pagos (solo obras a las que todavía les queda algo por pagar)
            if sum(o["_pagosRaw"]) > 0.5:
                lotes[o["clave"]]["desplazamiento"] = max(lotes[o["clave"]]["desplazamiento"], o["desplazamiento"])
    for rt in retenidos:
        if rt["enMapa"]:
            add_into(lotes[rt["clave"]]["egresosMes"], rt["_pagosRaw"])
    for l in lics:
        cat = categoria_licitacion(l)
        tag(l["clave"], cat)
        if l["enMapa"]:
            add_into(lotes[l["clave"]]["ingresosMes"], l["_cobrosRaw"])
            if cat == "vende":
                lotes[l["clave"]]["tramo"] = l["tramo"]
    for c in cartera:
        if c["situacion"] == "Con plan" and c["totalHorizonte"] > 0:
            tag(c["clave"], "cartera")
        elif c["situacion"].startswith("Sin plan"):
            tag(c["clave"], "legal")
        if c["enMapa"]:
            add_into(lotes[c["clave"]]["ingresosMes"], c["_cobrosRaw"])
            if c["_recupRaw"]:
                lotes[c["clave"]]["ingresosMes"][idx[c["mesRecupLegal"]]] += c["_recupRaw"]

    lotes_out = {}
    for clave, d in sorted(lotes.items()):
        cats = sorted(d["categorias"], key=CATEGORIAS.index)
        if not cats:
            continue  # solo retenidos sin obra: no colorea
        lotes_out[clave] = {
            "categoria": cats[0],
            "categorias": cats,
            "ingresos": r2(sum(d["ingresosMes"])),
            "egresos": r2(sum(d["egresosMes"])),
            "ingresosMes": [r2(v) for v in d["ingresosMes"]],
            "egresosMes": [r2(v) for v in d["egresosMes"]],
            "tramo": d["tramo"] if cats[0] == "vende" else 0,
            "desplazamiento": d["desplazamiento"] if cats[0] == "obra" else 0,
        }

    # cuadre de ingresos/egresos por lote contra Lotes_Maestro (sin los lotes de Santiago «por definir», que el
    # modelo carga en las filas de 01-01-01 y 02-01-01)
    genericas = {l["claveFuente"] for l in lics if l.get("claveFuente")}
    for row in maestro:
        clave = row["ClaveMapa"]
        if not RE_CLAVE_REAL.match(clave or "") or clave in genericas:
            continue  # claves virtuales: no van al mapa
        ing_m, egr_m = num(row["IngresosHorizonte"]), num(row["EgresosObrasHorizonte"])
        mine = lotes_out.get(clave, {"ingresos": 0.0, "egresos": 0.0})
        if abs(mine["ingresos"] - ing_m) > TOL or abs(mine["egresos"] - egr_m) > TOL:
            print(f"   aviso: {clave} JSON ing/egr {mine['ingresos']/1e6:,.2f}/{mine['egresos']/1e6:,.2f} "
                  f"vs Lotes_Maestro {ing_m/1e6:,.2f}/{egr_m/1e6:,.2f}")

    # --- totales por proyecto (gastos fijos no se reparten)
    por_proyecto = {}
    for p in sorted({o["proyecto"] for o in obras} | {l["proyecto"] for l in lics} | {c["proyecto"] for c in cartera}
                    | {rt["proyecto"] for rt in retenidos}):
        ing = defaultdict(lambda: zeros(n))
        egr = defaultdict(lambda: zeros(n))
        for c in cartera:
            if c["proyecto"] == p:
                add_into(ing["cartera"], c["_cobrosRaw"])
                if c["_recupRaw"]:
                    ing["recupLegal"][idx[c["mesRecupLegal"]]] += c["_recupRaw"]
        for l in lics:
            if l["proyecto"] == p:
                add_into(ing[GRUPO_CAMPO[l["grupo"]]], l["_cobrosRaw"])
        for o in obras:
            if o["proyecto"] == p:
                add_into(egr["mantGeneral" if o["tipo"].lower() == "reserva" else "obras"], o["_pagosRaw"])
        for rt in retenidos:
            if rt["proyecto"] == p:
                add_into(egr["retenidos"], rt["_pagosRaw"])
        ing_keys = ING_KEYS
        egr_keys = ["obras", "mantGeneral", "retenidos"]
        ing_d = {k: [r2(v) for v in ing[k]] for k in ing_keys}
        egr_d = {k: [r2(v) for v in egr[k]] for k in egr_keys}
        ing_d["total"] = [r2(sum(v)) for v in zip(*(ing[k] for k in ing_keys))]
        egr_d["total"] = [r2(sum(v)) for v in zip(*(egr[k] for k in egr_keys))]
        por_proyecto[p] = {"ingresos": ing_d, "egresos": egr_d}

    strip = lambda rows: [{k: v for k, v in d.items() if not k.startswith("_")} for d in rows]  # noqa: E731
    R = lambda arr: [r2(v) for v in arr]  # noqa: E731
    esc = {
        "kpis": kpis,
        "kpis2027": kpis2027,
        "mensual": {
            # la página trata «otros» como el ingreso que no viene de lotes: incluye las inversiones (v1.13+)
            "ingresos": {k: R(f[k]) for k in ING_KEYS}
            | {"otros": R([a + b for a, b in zip(f["otros"], f["inversiones"])]), "inversiones": R(f["inversiones"])}
            | {"total": R(f["totalIngresos"])},
            "egresos": {"obrasCJB": R(f["obrasCJB"]), "obrasCPAG": R(f["obrasCPAG"]), "mantGeneral": R(f["mantGeneral"]),
                        "complementarias": R(f["complementarias"]), "costoAtraso": R(f["costoAtraso"]),
                        "retenidos": R(f["retenidos"]), "gastosFijos": R(gastos_fijos),
                        "gastosFijosDetalle": {"rrhh": R(f["gfRRHH"]), "cargaFija": R(f["gfCargaFija"]),
                                               "admin": R(f["gfAdmin"]), "mantenimientos": R(f["gfMant"]),
                                               "itbis": R(f["gfITBIS"]), "provision": R(f["gfProvision"])},
                        "obligaciones": R(f["obligaciones"]), "pacc": R(f["gfBienes"]),
                        "total": R(f["totalEgresos"])},
            "saldoInicial": R(f["saldoInicial"]),
            "neto": R(f["neto"]),
            "lineaCredito": {"desembolso": R(f["lineaDesembolso"]), "pago": R(f["lineaPago"]),
                             "intereses": R(f["lineaIntereses"]), "deuda": R(f["lineaDeuda"])},
            "saldo": R(f["saldo"]),
            "colchon": R(f["colchon"]),
            "estado": f["estado"],
        },
        "porProyecto": por_proyecto,
        "obras": strip(obras),
        "licitaciones": strip(lics),
        "cartera": strip(cartera),
        "retenidos": strip(retenidos),
        "lotes": lotes_out,
        # v1.15+: supuestos de la combinación, regla 75/25, techos de compras y lo que muestra el Panel
        "params": params,
        "split": leer_split(wb),
        "techos": leer_techos(wb),
        "panel": leer_panel_extra(wb),
        "compras": leer_compras_combo(wb),
    }
    cuadrar_bloques(esc)
    return nombre, esc, maestro


def cuadrar_bloques(esc):
    """Lo que la página muestra de compras, Santiago y 75/25 debe ser lo mismo en el Panel y en sus hojas de cálculo."""
    sp, te, pa = esc["split"], esc["techos"], esc["panel"]
    pares = []
    if sp and pa:
        pares += [("techo global (Panel / Split_Santiago)", pa.get("techoCompras"), sp["techoGlobal"]),
                  ("capacidad de caja (Panel / Split_Santiago)", pa.get("capacidadCompras"), sp["disponible"]),
                  ("reserva Santiago (Panel / Split_Santiago)", pa.get("reservaSantiago"), sp["reservaSantiago"])]
        if pa.get("ratioNoObra") is not None and abs(pa["ratioNoObra"] - sp["ratio"]) > 1e-6:
            fail(f"No cuadra no-obra/obra: Panel {pa['ratioNoObra']} vs Split_Santiago {sp['ratio']}")
    if te and sp:
        pares += [("techo global (Techos_Areas / Split_Santiago)", te["global"], sp["techoGlobal"]),
                  ("capacidad (Techos_Areas / Split_Santiago)", te["capacidad"], sp["disponible"])]
    if te:
        pares.append(("techo asignado (Techos_Areas: tabla / fila «Techo asignado»)",
                      sum(a["techo"] for a in te["areas"]), te["asignado"]))
    for nombre, a, b in pares:
        if a is not None and abs(a - b) > TOL:
            fail(f"No cuadra {nombre}: {a:,.2f} vs {b:,.2f}")


def modelo_vigente():
    """Versión más alta del modelo (v1 < v1.1 < v1.2 …); a igual versión, el guardado más reciente."""
    def version(p):
        m = re.search(r" v(\d+(?:\.\d+)*)", p.name)
        return tuple(int(x) for x in m.group(1).split(".")) if m else (0,)

    candidatos = [p for p in MODELO_DIR.glob(MODELO_PATRON) if not p.name.startswith("~$")]
    if not candidatos:
        fail(f"No encontré «{MODELO_PATRON}» en {MODELO_DIR}")
    return max(candidatos, key=lambda p: (version(p), p.stat().st_mtime))


def comparar_hoja_versiones(data, hoja):
    """Compara la Base de cada versión con los valores fijos de la hoja Versiones (calculados con la Base cuando se armó
    esa hoja). Una diferencia no detiene la exportación: indica que algún supuesto cambió después."""
    meses = data["meses"]
    i27 = [i for i, m in enumerate(meses) if m.startswith("2027")]
    out = []
    for n, d in sorted(hoja.items()):
        esc = data["combos"].get(f"{n}|Base")
        if not esc or not i27:
            continue
        M, sp = esc["mensual"], esc.get("split") or {}
        k27 = esc.get("kpis2027") or esc["kpis"]
        calc = {"ingresos2027": sum(M["ingresos"]["total"][i] for i in i27),
                "egresos2027": sum(M["egresos"]["total"][i] for i in i27),
                "saldoMin2027": min(M["saldo"][i] for i in i27), "saldoDic27": k27["saldoFinal"],
                "saldoDic28": esc["kpis"]["saldoFinal"] if meses[-1] == "2028-12" else None,
                "capacidad": sp.get("disponible"), "techo": sp.get("techoGlobal"),
                "santiago": sp.get("reservaSantiago"), "ratio": sp.get("ratio")}
        for k, fijo in d["fijos"].items():
            if not isinstance(fijo, (int, float)) or calc.get(k) is None:
                continue
            es_ratio = k == "ratio"
            out.append({"version": n, "concepto": k, "hoja": round(fijo, 6) if es_ratio else r2(fijo),
                        "exportado": round(calc[k], 6) if es_ratio else r2(calc[k]),
                        "ok": abs(calc[k] - fijo) <= (1e-4 if es_ratio else TOL)})
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--src", help="Excel del modelo (por defecto, la versión más alta en la carpeta 2027)")
    ap.add_argument("--out", default=str(DEFAULT_OUT))
    ap.add_argument("--solo-activa", action="store_true",
                    help="Solo la versión elegida en el Panel D7 (3 recálculos en vez de 24)")
    args = ap.parse_args()
    src = Path(args.src) if args.src else modelo_vigente()
    if not src.exists():
        fail(f"No existe el Excel: {src}")
    print(f"Modelo: {src.name}")

    t0 = time.time()
    with tempfile.TemporaryDirectory(prefix="vbc_flujo_") as tmp:
        copia = copiar_modelo(src, tmp)
        # Lo que no cambia con la versión ni con el escenario se lee de la copia tal como la guardó el usuario
        wb0 = openpyxl.load_workbook(copia, data_only=True)
        versiones, reprog = leer_versiones(wb0)
        hoja_versiones = leer_hoja_versiones(wb0)
        etiqueta_activa = valor_nombre(wb0, "VersionActiva")
        n_activa = next((v["n"] for v in versiones if v["etiqueta"] == etiqueta_activa), 0)
        if versiones and not n_activa:
            fail(f"Panel D7 «{etiqueta_activa}» no está en la lista de versiones de Decisiones (sección 8)")
        for v in versiones:
            v.update(hoja_versiones.get(v["n"], {}).get("desc", {}))
        data = {
            "esquema": 2,
            "generado": dt.datetime.now().isoformat(timespec="seconds"),
            "fuente": src.name,
            "fuenteModificado": dt.datetime.fromtimestamp(src.stat().st_mtime).isoformat(timespec="seconds"),
            "unidad": "RD$",
            "meses": None,
            # lo que estaba elegido en el Excel guardado: la página abre con esta versión y este escenario
            "activa": {"version": n_activa, "escenario": wb0["Panel"]["D5"].value,
                       "fuenteCubicaciones": valor_nombre(wb0, "FuenteCub", ""),
                       "ret1Remitida": valor_nombre(wb0, "Ret1Remitida", "")},
            "versiones": versiones,
            "reprogramacion": reprog,
            "riesgos": leer_riesgos(wb0),
            "compras": leer_compras_hist(wb0),
            "gastoReal": leer_gasto_real(wb0),
            "combos": {},
            "maestro": {},
        }
        wb0.close()
        a_recalcular = [v for v in versiones if not args.solo_activa or v["n"] == n_activa]
        copias = recalcular(copia, tmp, [(v["n"], v["etiqueta"]) for v in a_recalcular])
        print()
        print(f"{'Versión':<8}{'Escenario':<11}{'Ingresos':>9}{'Egresos':>9}{'Dic-27':>8}{'Dic-28':>8}{'Mínimo':>8}"
              f"  Mes min  {'Caja compras':>12}{'Techo':>6}{'Stgo':>7}  No-obra  (RD$ M)")
        for (n, e), p in copias.items():
            wb = openpyxl.load_workbook(p, data_only=True)
            if data["meses"] is None:
                data["meses"] = meses_flujo(wb["Flujo_Mensual"])
            nombre, esc, maestro = armar_escenario(wb, data["meses"])
            if nombre != e:
                fail(f"Panel!D5 dice «{nombre}» en la copia de «{e}»")
            if n and int(num(esc["params"].get("versionN"))) != n:
                fail(f"La copia de la versión {n} ({e}) tiene VersionN = {esc['params'].get('versionN')}")
            data["combos"][f"{n}|{e}"] = esc
            M = 1e6
            k, k27, sp = esc["kpis"], esc["kpis2027"] or esc["kpis"], esc["split"] or {}
            print(f"{n:<8}{e:<11}{k['ingresos']/M:>9,.1f}{k['egresos']/M:>9,.1f}{k27['saldoFinal']/M:>8,.1f}"
                  f"{k['saldoFinal']/M:>8,.1f}{k['saldoMin']/M:>8,.1f}  {k['mesSaldoMin']}  "
                  f"{num(sp.get('disponible'))/M:>12,.1f}{num(sp.get('techoGlobal'))/M:>6,.1f}"
                  f"{num(sp.get('reservaSantiago'))/M:>7,.1f}  {num(sp.get('ratio')):.1%}")
            if not data["maestro"]:
                for row in maestro:
                    data["maestro"][row["ClaveMapa"]] = {
                        "proyecto": row["Proyecto"], "id": row["IDLote"], "tipo": row["Tipo"] or "",
                        "subtipo": row["Subtipo"] or "", "nombre": row["Nombre"] or "",
                        "desarrollador": row["Desarrollador"] or "", "area": r2(num(row["AreaM2"])),
                        "estado": row["Estado"] or "",
                    }
            wb.close()

    data["controlVersiones"] = comparar_hoja_versiones(data, hoja_versiones)
    malos = [x for x in data["controlVersiones"] if not x["ok"]]
    if malos:
        print(f"\nHoja Versiones (valores fijos de la Base) vs exportación: {len(malos)} diferencias "
              "(algún supuesto cambió después de calcular esa hoja):")
        for x in malos:
            print(f"   versión {x['version']} · {x['concepto']}: hoja {x['hoja']:,} · exportado {x['exportado']:,}")
    elif data["controlVersiones"]:
        print(f"\nHoja Versiones: las {len(data['controlVersiones'])} cifras fijas de la Base coinciden con la exportación.")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"\nEscrito {out}  ({out.stat().st_size/1024:,.0f} KB) en {time.time() - t0:,.0f} s")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
