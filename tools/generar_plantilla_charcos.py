"""
Genera la plantilla de Excel con la que la Dirección Técnica levanta los charcos de Ciudad Juan Bosch.

Una fila por charco. Lo esencial:
  · Ubicación: se pega tal cual la ubicación de Google Maps (enlace, coordenadas o código plus);
    el generador del mapa la convierte en coordenadas.
  · Lote o calle donde está (lista con ID y nombre) y un sub-nombre libre.
  · Criticidad (rojo / amarillo / verde) y de dónde viene el agua (calle o lote de origen).
La columna «Revisión» avisa en la misma fila qué dato obligatorio falta.

El archivo lleno se guarda como Activos/ANALISIS_CHARCOS_CJB.xlsx y se carga en el mapa con
    python tools/generar_activos_cjb.py

Uso:
    python tools/generar_plantilla_charcos.py
Salida:
    Activos/PLANTILLA_ANALISIS_CHARCOS_CJB.xlsx
"""

import csv
import datetime as dt
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_PATH = os.path.join(ROOT, "Datos de CJB.csv")
SALIDA = os.path.join(ROOT, "Activos", "PLANTILLA_ANALISIS_CHARCOS_CJB.xlsx")

F, FS = "Segoe UI", "Segoe UI Semibold"
NAVY, RED, SLATE, INK, LINE, SOFT = "1B3A6B", "E63946", "64748B", "0F172A", "E2E8F0", "F1F5F9"
ENTIDAD = "FIDEICOMISO PARA LA CONSTRUCCIÓN DE VIVIENDAS DE BAJO COSTO (FIDEICOMISO VBC RD)"
N_FILAS = 200
P = 6  # primera fila de datos (fila 4: pista · fila 5: encabezado)

CRITICIDAD = [
    ("Rojo – Muy crítico", "FEE2E2", "991B1B",
     "Impide el paso de vehículos o peatones, entra a viviendas o locales, o dura muchas horas."),
    ("Amarillo – Intermedio", "FEF3C7", "92400E",
     "Dificulta el paso y molesta a los residentes, pero se seca en pocas horas."),
    ("Verde – Menor", "DCFCE7", "166534",
     "Acumulación pequeña o momentánea que no impide el paso."),
]
ORIGEN = [
    "Lluvia que corre por la calle",
    "Sale de un lote habitacional",
    "Sale de un lote comercial",
    "Sale de un lote institucional o equipamiento",
    "Sale de una obra en construcción",
    "De varios lugares (detallar)",
    "No se sabe",
]
CAUSAS = [
    "Falta imbornal o filtrante",
    "Imbornal obstruido (basura, lodo, arena)",
    "Imbornal insuficiente para el agua que llega",
    "Desnivel o pendiente de la calle",
    "Agua que sale de un lote u obra",
    "Tubería rota o fuga",
    "Otra (detallar en observaciones)",
]
AFECTA = ["Paso de vehículos", "Paso de peatones", "Viviendas o locales", "Acera o área común", "Nada importante"]

# (encabezado, obligatoria, ancho, tipo, lista, pista)
COLUMNAS = [
    ("ID del charco", False, 10, "auto", None, "Ya numerado"),
    ("Revisión", False, 24, "revision", None, "Se llena sola"),
    ("Fecha del levantamiento", True, 13, "fecha", None, "dd/mm/aaaa"),
    ("Ubicación en Google Maps", True, 34, "texto", None, "Pegue el enlace o las coordenadas"),
    ("Calle o punto de referencia", True, 30, "texto", None, "Ej.: Av. La Gaviota frente a la escuela"),
    ("Lote o calle donde está", False, 36, "lista", "lst_lotes", "Escriba parte del nombre o el ID"),
    ("Sub-nombre o detalle del lugar", False, 24, "texto", None, "Ej.: Torre 4, parqueo, esquina"),
    ("Criticidad", True, 19, "lista", "lst_criticidad", "Rojo · Amarillo · Verde"),
    ("¿De dónde viene el agua?", True, 30, "lista", "lst_origen", "Seleccione"),
    ("Lote de donde sale el agua", False, 36, "lista", "lst_lotes", "Si sale de un lote"),
    ("Otros lotes de donde sale el agua", False, 20, "texto", None, "IDs separados por coma"),
    ("Causa probable", False, 30, "lista", "lst_causas", "Seleccione"),
    ("¿Qué afecta?", False, 20, "lista", "lst_afecta", "Seleccione"),
    ("Tamaño aproximado (m²)", False, 11, "numero", None, "Número"),
    ("Profundidad aproximada (cm)", False, 12, "numero", None, "Número"),
    ("Tiempo en secarse (horas)", False, 11, "numero", None, "Número"),
    ("Fotos o video (enlace)", False, 26, "texto", None, "OneDrive, Drive o WhatsApp"),
    ("Observaciones", False, 38, "texto", None, "Texto libre"),
]
SIN_AJUSTE = {"Ubicación en Google Maps", "Fotos o video (enlace)"}   # enlaces largos: una sola línea
COL = {c[0]: get_column_letter(i) for i, c in enumerate(COLUMNAS, start=1)}

EJEMPLO = {
    "ID del charco": "CH-001",
    "Revisión": "✔ Completa",
    "Fecha del levantamiento": "28/10/2026",
    "Ubicación en Google Maps": "18.491523, -69.745211   (o https://maps.app.goo.gl/…)",
    "Calle o punto de referencia": "Calle Rey David esquina Cayo Confites",
    "Lote o calle donde está": "15-01-03 · Calle Rey David",
    "Sub-nombre o detalle del lugar": "Frente a la parada",
    "Criticidad": "Rojo – Muy crítico",
    "¿De dónde viene el agua?": "Sale de un lote habitacional",
    "Lote de donde sale el agua": "04-11-01 · (nombre del lote)",
    "Causa probable": "Imbornal obstruido (basura, lodo, arena)",
    "¿Qué afecta?": "Paso de vehículos",
    "Tamaño aproximado (m²)": 40,
    "Profundidad aproximada (cm)": 15,
    "Tiempo en secarse (horas)": 6,
}


def fill(c):
    return PatternFill("solid", fgColor=c, bgColor=c)


def font(sz=10, color=INK, name=F, bold=False, italic=False):
    return Font(name=name, size=sz, color=color, bold=bold, italic=italic)


def banda(ws, ult, titulo, sub):
    for r in (1, 2):
        for c in range(1, ult + 1):
            ws.cell(row=r, column=c).fill = fill(NAVY)
    for c in range(1, ult + 1):
        ws.cell(row=3, column=c).fill = fill(RED)
    ws["A1"], ws["A2"] = titulo, sub
    ws["A1"].font = font(15, "FFFFFF", FS)
    ws["A2"].font = font(9, "C7D2FE")
    ws["A1"].alignment = Alignment(vertical="bottom", indent=1)
    ws["A2"].alignment = Alignment(vertical="top", indent=1)
    ws.row_dimensions[1].height, ws.row_dimensions[2].height, ws.row_dimensions[3].height = 28, 20, 3
    ws.sheet_view.showGridLines = False


def opcion_lote(l):
    nombre = (l.get("nombre") or "").strip() or (l.get("subtipo") or "").strip() or l.get("tipo", "")
    return f"{l['id']} · {nombre}"


def formula_revision(r):
    c = lambda k: f"{COL[k]}{r}"
    fecha, ubic, ref = c("Fecha del levantamiento"), c("Ubicación en Google Maps"), c("Calle o punto de referencia")
    crit, ori, lote_o = c("Criticidad"), c("¿De dónde viene el agua?"), c("Lote de donde sale el agua")
    otros = c("Otros lotes de donde sale el agua")
    rango = f"{COL['Fecha del levantamiento']}{r}:{COL['Observaciones']}{r}"
    sale = f'LEFT({ori},7)="Sale de"'
    falta_lote = f'AND({sale},{lote_o}="",{otros}="")'
    return (f'=IF(COUNTA({rango})=0,"",IF(AND({fecha}<>"",{ubic}<>"",{ref}<>"",{crit}<>"",{ori}<>"",NOT({falta_lote})),'
            f'"✔ Completa","Falta: "&IF({fecha}="","fecha · ","")&IF({ubic}="","ubicación · ","")'
            f'&IF({ref}="","referencia · ","")&IF({crit}="","criticidad · ","")&IF({ori}="","origen del agua · ","")'
            f'&IF({falta_lote},"lote de origen · ","")))')


def hoja_instrucciones(wb, lotes):
    por_id = {l["id"]: l for l in lotes}
    ejemplo = dict(EJEMPLO)
    for k in ("Lote o calle donde está", "Lote de donde sale el agua"):
        lid = ejemplo[k].split(" ")[0]
        if lid in por_id:
            ejemplo[k] = opcion_lote(por_id[lid])
    ins = wb.create_sheet("Instrucciones", 0)
    banda(ins, 4, "Levantamiento de charcos · Ciudad Juan Bosch",
          f"Cómo llenar la plantilla · {ENTIDAD}")
    ins.column_dimensions["A"].width = 3
    ins.column_dimensions["B"].width = 28
    ins.column_dimensions["C"].width = 96
    r = 5

    def titulo(t):
        nonlocal r
        c = ins.cell(row=r, column=2, value=t)
        c.font = font(11.5, NAVY, FS)
        ins.row_dimensions[r].height = 22
        r += 1

    def par(etq, texto, alto=34):
        nonlocal r
        a = ins.cell(row=r, column=2, value=etq)
        a.font = font(10, INK, FS)
        a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ins.cell(row=r, column=3, value=texto)
        b.font = font(10)
        b.alignment = Alignment(vertical="top", wrap_text=True)
        ins.row_dimensions[r].height = alto
        r += 1

    titulo("1. Una fila por charco")
    par("Hoja «Charcos»",
        "Cada charco o punto donde se acumula el agua va en una fila. El ID ya está numerado (CH-001, CH-002…). "
        "Las columnas con encabezado azul son obligatorias; la columna «Revisión» le dice en la misma fila "
        "qué falta. Cuando diga «✔ Completa», la fila está lista.", 48)
    r += 1
    titulo("2. Ubicación desde Google Maps (lo más importante)")
    par("En el celular",
        "1) Abra Google Maps y acérquese al charco.  2) Mantenga presionado el punto exacto hasta que aparezca "
        "el pin rojo (si está parado junto al charco, mantenga presionado sobre su punto azul).  "
        "3) Toque las coordenadas que salen arriba para copiarlas, o use «Compartir» → «Copiar enlace».  "
        "4) Péguelo en la columna «Ubicación en Google Maps».", 62)
    par("En la computadora",
        "Clic derecho sobre el punto del charco → clic en las coordenadas (por ejemplo 18.491523, -69.745211); "
        "quedan copiadas. Péguelas en la columna «Ubicación en Google Maps».", 34)
    par("Qué se acepta",
        "Coordenadas (18.491523, -69.745211), enlace de Google Maps (https://maps.app.goo.gl/… o "
        "https://www.google.com/maps/…), código plus (FHRW+J2) o coordenadas UTM. Nosotros obtenemos la "
        "ubicación exacta a partir de lo que pegue.", 46)
    r += 1
    titulo("3. Lote o calle")
    par("Lote o calle donde está",
        "Despliegue la lista y escriba parte del nombre o del ID (por ejemplo «Rey David» o «15-01»). La lista "
        "trae todos los lotes y calles del Master Plan; la hoja «Lotes» sirve de consulta. Si el lugar no "
        "aparece o quiere precisar (torre, parqueo, esquina), escríbalo en «Sub-nombre o detalle del lugar».", 48)
    par("¿De dónde viene el agua?",
        "Indique si es lluvia que corre por la calle o si sale de un lote. Si sale de un lote, elija cuál en "
        "«Lote de donde sale el agua»; si son varios, escriba los demás IDs en «Otros lotes…», separados por "
        "coma. En el mapa se dibuja el recorrido del agua desde esos lotes hasta el charco.", 48)
    r += 1
    titulo("4. Criticidad")
    for nombre, bg, fg, desc in CRITICIDAD:
        a = ins.cell(row=r, column=2, value=nombre)
        a.font = font(10, fg, FS)
        a.fill = fill(bg)
        a.alignment = Alignment(vertical="center", indent=1)
        b = ins.cell(row=r, column=3, value=desc)
        b.font = font(10)
        b.alignment = Alignment(vertical="center", wrap_text=True)
        ins.row_dimensions[r].height = 24
        r += 1
    r += 1
    titulo("5. Ejemplo de una fila llena")
    for k, v in ejemplo.items():
        par(k, str(v), 18)
    r += 1
    titulo("6. Al terminar")
    par("Envío",
        "Guarde el archivo agregando su nombre y la fecha (por ejemplo «Charcos CJB - Delvinson - 2026-11-05») y "
        "envíelo al equipo de Data del Fideicomiso. Puede enviarlo por partes: se cargan todas las filas "
        "que estén completas.", 46)
    ins.sheet_view.zoomScale = 100
    ins.sheet_view.tabSelected = False
    return ins


def main():
    with open(CSV_PATH, encoding="utf-8-sig") as f:
        lotes = sorted(csv.DictReader(f), key=lambda x: x["id"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Charcos"
    hoja_instrucciones(wb, lotes)

    # ---------------------------------------------------------------- Charcos
    ult = len(COLUMNAS)
    banda(ws, ult, "Levantamiento de charcos · Ciudad Juan Bosch",
          f"Una fila por charco · encabezado azul = obligatorio · vea la hoja «Instrucciones» · {ENTIDAD}")
    fin = P + N_FILAS - 1
    for j, (enc, obl, ancho, tipo, lista, pista) in enumerate(COLUMNAS, start=1):
        L = get_column_letter(j)
        ws.column_dimensions[L].width = ancho
        p = ws.cell(row=4, column=j, value=pista)
        p.font = font(8, "94A3B8", italic=True)
        p.alignment = Alignment(horizontal="center", wrap_text=True)
        h = ws.cell(row=5, column=j, value=enc + (" *" if obl else ""))
        h.font = font(9.5, "FFFFFF", FS)
        h.fill = fill(NAVY if obl else SLATE)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = Border(left=Side(style="thin", color="FFFFFF"), right=Side(style="thin", color="FFFFFF"))
        for i in range(N_FILAS):
            r = P + i
            c = ws.cell(row=r, column=j)
            c.font = font(9.5)
            c.border = Border(bottom=Side(style="thin", color=LINE), right=Side(style="thin", color=LINE))
            c.alignment = Alignment(vertical="center", wrap_text=tipo == "texto" and enc not in SIN_AJUSTE,
                                    horizontal="center" if tipo in ("auto", "numero", "fecha") else "left")
            if tipo == "auto":
                c.value = f"CH-{i + 1:03d}"
                c.font = font(9.5, SLATE, FS)
                c.fill = fill(SOFT)
            elif tipo == "revision":
                c.value = formula_revision(r)
                c.font = font(9, SLATE)
                c.fill = fill(SOFT)
            else:
                c.protection = Protection(locked=False)
                if tipo == "texto":
                    c.number_format = "@"
                elif tipo == "fecha":
                    c.number_format = "dd/mm/yyyy"
        rango = f"{L}{P}:{L}{fin}"
        dv = None
        if tipo == "lista":
            dv = DataValidation(type="list", formula1=f"={lista}", allow_blank=True, showErrorMessage=True,
                                errorStyle="warning", errorTitle="Valor fuera de la lista",
                                error="Ese valor no está en la lista. ¿Desea dejarlo así?")
        elif tipo == "numero":
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                showErrorMessage=True, errorTitle="Valor no válido", error="Escriba solo un número.")
        elif tipo == "fecha":
            dv = DataValidation(type="date", operator="greaterThanOrEqual", formula1="DATE(2025,1,1)",
                                allow_blank=True, showErrorMessage=True, errorTitle="Fecha no válida",
                                error="Escriba la fecha como dd/mm/aaaa.")
        if dv:
            dv.add(rango)
            ws.add_data_validation(dv)
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 38
    for i in range(N_FILAS):
        ws.row_dimensions[P + i].height = 20

    cr, rv = COL["Criticidad"], COL["Revisión"]
    for nombre, bg, fg, _ in CRITICIDAD:
        ws.conditional_formatting.add(f"{cr}{P}:{cr}{fin}", FormulaRule(
            formula=[f'LEFT(${cr}{P},4)="{nombre[:4]}"'], fill=fill(bg), font=Font(color=fg, bold=True)))
    ws.conditional_formatting.add(f"{rv}{P}:{rv}{fin}", FormulaRule(
        formula=[f'LEFT(${rv}{P},1)="✔"'], fill=fill("DCFCE7"), font=Font(color="166534", bold=True)))
    ws.conditional_formatting.add(f"{rv}{P}:{rv}{fin}", FormulaRule(
        formula=[f'LEFT(${rv}{P},5)="Falta"'], fill=fill("FFEDD5"), font=Font(color="9A3412")))
    ws.freeze_panes = f"C{P}"
    ws.auto_filter.ref = f"A5:{get_column_letter(ult)}{fin}"
    pr = ws.protection
    pr.sheet = True
    pr.autoFilter = False
    pr.formatColumns = False
    pr.formatRows = False
    pr.formatCells = False

    # ---------------------------------------------------------------- Lotes (consulta)
    lo = wb.create_sheet("Lotes")
    banda(lo, 5, "Lotes y calles del Master Plan", "Consulta para las columnas de lote o calle")
    for j, (t, w) in enumerate([("ID", 11), ("Nombre", 50), ("Tipo", 16), ("Subtipo", 26),
                                ("Como aparece en la lista", 56)], start=1):
        c = lo.cell(row=5, column=j, value=t)
        c.font = font(9.5, "FFFFFF", FS)
        c.fill = fill(NAVY)
        lo.column_dimensions[get_column_letter(j)].width = w
    for i, l in enumerate(lotes):
        for j, v in enumerate([l["id"], l.get("nombre") or None, l.get("tipo") or None, l.get("subtipo") or None,
                               opcion_lote(l)], start=1):
            c = lo.cell(row=6 + i, column=j, value=v)
            c.font = font(9.5, SLATE if j == 5 else INK)
            c.border = Border(bottom=Side(style="thin", color=LINE))
    lo.freeze_panes = "B6"
    lo.auto_filter.ref = f"A5:E{5 + len(lotes)}"
    wb.defined_names["lst_lotes"] = DefinedName("lst_lotes", attr_text=f"Lotes!$E$6:$E${5 + len(lotes)}")

    # ---------------------------------------------------------------- Listas (oculta)
    ls = wb.create_sheet("Listas")
    listas = [("lst_criticidad", [c[0] for c in CRITICIDAD]), ("lst_origen", ORIGEN),
              ("lst_causas", CAUSAS), ("lst_afecta", AFECTA)]
    for j, (nombre, vals) in enumerate(listas, start=1):
        L = get_column_letter(j)
        ls.cell(row=1, column=j, value=nombre)
        for i, v in enumerate(vals, start=2):
            ls.cell(row=i, column=j, value=v)
        wb.defined_names[nombre] = DefinedName(nombre, attr_text=f"Listas!${L}$2:${L}${len(vals) + 1}")
    ls.sheet_state = "hidden"

    for hoja in wb.worksheets:
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 1          # abre en «Charcos»
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "Levantamiento de charcos · Ciudad Juan Bosch"
    wb.properties.creator = "Fideicomiso VBC RD"
    wb.properties.created = dt.datetime.now()
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    wb.save(SALIDA)
    print(f"Plantilla: {SALIDA}  ({N_FILAS} filas · {len(lotes)} lotes y calles en la lista)")


if __name__ == "__main__":
    main()
