"""
Script para generar los archivos de Excel para los mapas CJB y CPAG.
Entregables:
1. LISTADO_LOTES_MAPA_CJB.xlsx (Ciudad Juan Bosch - 213 lotes)
2. LISTADO_LOTES_MAPA_CPAG.xlsx (Ciudad Presidente Antonio Guzmán - 98 lotes)

Columnas requeridas por el usuario:
- ID
- Nombre Comercial
- Descripción
- Estatus
Incluye además hoja con detalle técnico completo y resumen ejecutivo.
"""

import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Estilos corporativos ejecutivos
FONT_FAMILY = "Segoe UI"
COLOR_NAVY = "1E3A8A"       # Azul institucional oscuro
COLOR_NAVY_LIGHT = "2563EB" # Azul acento
COLOR_TEXT_MAIN = "0F172A"  # Slate oscuro
COLOR_WHITE = "FFFFFF"
COLOR_ZEBRA = "F8FAFC"      # Gris tenue alternado
BORDER_LIGHT = "CBD5E1"     # Gris claro bordes

# Colores de estatus
STATUS_COLORS = {
    'Terminado': {'fill': 'DCFCE7', 'font': '166534'},
    'Concluido': {'fill': 'DCFCE7', 'font': '166534'},
    'En Construcción': {'fill': 'FEF3C7', 'font': '92400E'},
    'En Construccion': {'fill': 'FEF3C7', 'font': '92400E'},
    'En Terminación': {'fill': 'FEF3C7', 'font': '92400E'},
    'Paralizado': {'fill': 'FEE2E2', 'font': '991B1B'},
    'En Licitación': {'fill': 'DBEAFE', 'font': '1E40AF'},
    'En Licitacion': {'fill': 'DBEAFE', 'font': '1E40AF'},
    'Adjudicado': {'fill': 'DBEAFE', 'font': '1E40AF'},
    'Disponible': {'fill': 'E0E7FF', 'font': '3730A3'},
    'No Iniciado': {'fill': 'F1F5F9', 'font': '475569'},
}

def get_border(thin_color=BORDER_LIGHT):
    thin = Side(border_style="thin", color=thin_color)
    return Border(left=thin, right=thin, top=thin, bottom=thin)

def get_total_border():
    top_thin = Side(border_style="thin", color="1E3A8A")
    bottom_double = Side(border_style="double", color="1E3A8A")
    return Border(top=top_thin, bottom=bottom_double)

def apply_title_header(ws, title_main, subtitle, universe_str):
    ws.merge_cells('A1:E1')
    ws['A1'] = "FIDEICOMISO PARA LA CONSTRUCCIÓN DE VIVIENDAS DE BAJO COSTO (FIDEICOMISO VBC RD)"
    ws['A1'].font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY)
    ws['A1'].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells('A2:E2')
    ws['A2'] = title_main
    ws['A2'].font = Font(name=FONT_FAMILY, size=14, bold=True, color=COLOR_NAVY)
    ws['A2'].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells('A3:E3')
    ws['A3'] = f"{subtitle} | {universe_str} | Base Cartográfica Oficial Sistema VBC"
    ws['A3'].font = Font(name=FONT_FAMILY, size=9, italic=True, color="64748B")
    ws['A3'].alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 10

def auto_adjust_columns(ws, max_cols=None):
    ws.views.sheetView[0].showGridLines = True
    for col in ws.columns:
        if max_cols and col[0].column > max_cols:
            continue
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip merged header rows 1-3 when calculating width
            if cell.row in [1, 2, 3]:
                continue
            if cell.value is not None:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def build_cjb_excel(dest_path):
    print("Building CJB Excel...")
    # Leer datos oficiales de CJB
    csv_path = os.path.join(ROOT, "Datos de CJB.csv")
    xlsx_master = os.path.join(ROOT, "MATRIZ_GENERAL_TODOS_LOS_LOTES_CJB.xlsx")

    # Usamos la matriz general limpia que tiene acentos perfectos y 'Por Definir'
    df_master = pd.read_excel(xlsx_master, sheet_name="Matriz General", header=4).dropna(subset=['Lote ID'])
    
    wb = openpyxl.Workbook()
    # Sheet 1: Listado de Lotes con las 4 columnas solicitadas
    ws1 = wb.active
    ws1.title = "Listado de Lotes"

    apply_title_header(
        ws1,
        "MAPA CJB - LISTADO OFICIAL DE LOTES (CIUDAD JUAN BOSCH)",
        "Listado con ID, Nombre Comercial, Descripción y Estatus",
        "Universo: 213 Lotes"
    )

    headers_ws1 = ["No.", "ID Lote", "Nombre Comercial", "Descripción", "Estatus"]
    ws1.row_dimensions[5].height = 26

    for col_idx, h in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=5, column=col_idx, value=h)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_WHITE)
        cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = get_border()

    border_thin = get_border()
    row_start = 6

    for i, (_, r) in enumerate(df_master.iterrows(), 1):
        curr_row = row_start + i - 1
        ws1.row_dimensions[curr_row].height = 20
        is_zebra = (i % 2 == 0)
        row_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_zebra else PatternFill(fill_type=None)

        lote_id = str(r["Lote ID"]).strip()
        nombre_comercial = str(r["Nombre del Proyecto / Denominación"]).strip()
        subtipo = str(r["Subtipo / Destino"]).strip()
        tipo = str(r["Tipo Suelo"]).strip()
        estatus = str(r["Estatus de Obra"]).strip()

        # Descripción: Combinar tipo y subtipo de manera clara y rica
        if subtipo and subtipo.lower() != tipo.lower() and subtipo != "Pendiente":
            desc = f"{tipo} - {subtipo}"
        else:
            desc = f"Lote {tipo}" if not subtipo or subtipo == tipo else f"{tipo} ({subtipo})"

        # 1. No.
        c1 = ws1.cell(row=curr_row, column=1, value=i)
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_TEXT_MAIN)
        c1.fill = row_fill
        c1.border = border_thin

        # 2. ID Lote (como texto estricto para evitar conversión a fecha)
        c2 = ws1.cell(row=curr_row, column=2, value=lote_id)
        c2.number_format = "@"
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY)
        c2.fill = row_fill
        c2.border = border_thin

        # 3. Nombre Comercial
        c3 = ws1.cell(row=curr_row, column=3, value=nombre_comercial)
        c3.alignment = Alignment(horizontal="left", vertical="center")
        c3.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_TEXT_MAIN)
        c3.fill = row_fill
        c3.border = border_thin

        # 4. Descripción
        c4 = ws1.cell(row=curr_row, column=4, value=desc)
        c4.alignment = Alignment(horizontal="left", vertical="center")
        c4.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_TEXT_MAIN)
        c4.fill = row_fill
        c4.border = border_thin

        # 5. Estatus
        c5 = ws1.cell(row=curr_row, column=5, value=estatus)
        c5.alignment = Alignment(horizontal="center", vertical="center")
        status_style = STATUS_COLORS.get(estatus, {'fill': 'F1F5F9', 'font': '475569'})
        c5.fill = PatternFill(start_color=status_style['fill'], end_color=status_style['fill'], fill_type="solid")
        c5.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=status_style['font'])
        c5.border = border_thin

    # Fila de totales
    total_row = row_start + len(df_master)
    ws1.row_dimensions[total_row].height = 22
    ws1.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    tot_lbl = ws1.cell(row=total_row, column=1, value="TOTAL LOTES REGISTRADOS EN MAPA CJB:")
    tot_lbl.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    tot_lbl.alignment = Alignment(horizontal="right", vertical="center")
    tot_lbl.border = get_total_border()

    tot_val = ws1.cell(row=total_row, column=5, value=len(df_master))
    tot_val.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    tot_val.alignment = Alignment(horizontal="center", vertical="center")
    tot_val.border = get_total_border()

    auto_adjust_columns(ws1, max_cols=5)
    ws1.freeze_panes = "A6"

    # Sheet 2: Inventario Técnico Completo
    ws2 = wb.create_sheet(title="Detalle Técnico Completo")
    apply_title_header(
        ws2,
        "MAPA CJB - INVENTARIO TÉCNICO DETALLADO (CIUDAD JUAN BOSCH)",
        "Base Técnica de Planificación, Superficie, Unidades y Desarrolladores",
        "Universo: 213 Lotes"
    )

    headers_ws2 = [
        "No.", "ID Lote", "Nombre Comercial", "Tipo de Suelo", "Subtipo / Destino",
        "Estatus de Obra", "Desarrollador / Entidad Responsable",
        "Superficie (m²)", "Viviendas Proyectadas", "Viviendas Concluidas", "Viviendas Pendientes"
    ]
    ws2.row_dimensions[5].height = 26
    for col_idx, h in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=5, column=col_idx, value=h)
        cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
        cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = get_border()

    for i, (_, r) in enumerate(df_master.iterrows(), 1):
        curr_row = row_start + i - 1
        ws2.row_dimensions[curr_row].height = 20
        is_zebra = (i % 2 == 0)
        row_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_zebra else PatternFill(fill_type=None)

        lote_id = str(r["Lote ID"]).strip()
        nombre_comercial = str(r["Nombre del Proyecto / Denominación"]).strip()
        tipo = str(r["Tipo Suelo"]).strip()
        subtipo = str(r["Subtipo / Destino"]).strip()
        estatus = str(r["Estatus de Obra"]).strip()
        desarrollador = str(r["Desarrollador / Entidad Responsable"]).strip() if pd.notnull(r["Desarrollador / Entidad Responsable"]) else "Sin Desarrollador"
        area = float(r["Superficie (m²)"]) if pd.notnull(r["Superficie (m²)"]) else 0.0
        u_proy = int(r["Unidades Proyectadas"]) if pd.notnull(r["Unidades Proyectadas"]) else 0
        u_conc = int(r["Unidades Concluidas"]) if pd.notnull(r["Unidades Concluidas"]) else 0
        u_pend = int(r["Unidades Pendientes"]) if pd.notnull(r["Unidades Pendientes"]) else 0

        ws2.cell(row=curr_row, column=1, value=i).alignment = Alignment(horizontal="center", vertical="center")
        
        c_id = ws2.cell(row=curr_row, column=2, value=lote_id)
        c_id.number_format = "@"
        c_id.alignment = Alignment(horizontal="center", vertical="center")
        c_id.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY)

        ws2.cell(row=curr_row, column=3, value=nombre_comercial).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(row=curr_row, column=4, value=tipo).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=curr_row, column=5, value=subtipo).alignment = Alignment(horizontal="left", vertical="center")

        c_st = ws2.cell(row=curr_row, column=6, value=estatus)
        c_st.alignment = Alignment(horizontal="center", vertical="center")
        status_style = STATUS_COLORS.get(estatus, {'fill': 'F1F5F9', 'font': '475569'})
        c_st.fill = PatternFill(start_color=status_style['fill'], end_color=status_style['fill'], fill_type="solid")
        c_st.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=status_style['font'])

        ws2.cell(row=curr_row, column=7, value=desarrollador).alignment = Alignment(horizontal="left", vertical="center")

        c_area = ws2.cell(row=curr_row, column=8, value=area)
        c_area.number_format = "#,##0.00"
        c_area.alignment = Alignment(horizontal="right", vertical="center")

        c_up = ws2.cell(row=curr_row, column=9, value=u_proy)
        c_up.number_format = "#,##0"
        c_up.alignment = Alignment(horizontal="right", vertical="center")

        c_uc = ws2.cell(row=curr_row, column=10, value=u_conc)
        c_uc.number_format = "#,##0"
        c_uc.alignment = Alignment(horizontal="right", vertical="center")

        c_un = ws2.cell(row=curr_row, column=11, value=u_pend)
        c_un.number_format = "#,##0"
        c_un.alignment = Alignment(horizontal="right", vertical="center")

        for c_idx in range(1, 12):
            cell = ws2.cell(row=curr_row, column=c_idx)
            cell.border = border_thin
            if c_idx != 6:
                cell.fill = row_fill
                cell.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_TEXT_MAIN)

    # Totales hoja 2
    ws2.row_dimensions[total_row].height = 22
    ws2.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
    tot_lbl2 = ws2.cell(row=total_row, column=1, value="TOTALES GENERALES CJB:")
    tot_lbl2.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    tot_lbl2.alignment = Alignment(horizontal="right", vertical="center")
    tot_lbl2.border = get_total_border()

    c_tot_area = ws2.cell(row=total_row, column=8, value=f"=SUM(H6:H{total_row-1})")
    c_tot_area.number_format = "#,##0.00"
    c_tot_area.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    c_tot_area.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_area.border = get_total_border()

    c_tot_up = ws2.cell(row=total_row, column=9, value=f"=SUM(I6:I{total_row-1})")
    c_tot_up.number_format = "#,##0"
    c_tot_up.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    c_tot_up.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_up.border = get_total_border()

    c_tot_uc = ws2.cell(row=total_row, column=10, value=f"=SUM(J6:J{total_row-1})")
    c_tot_uc.number_format = "#,##0"
    c_tot_uc.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    c_tot_uc.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_uc.border = get_total_border()

    c_tot_un = ws2.cell(row=total_row, column=11, value=f"=SUM(K6:K{total_row-1})")
    c_tot_un.number_format = "#,##0"
    c_tot_un.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    c_tot_un.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_un.border = get_total_border()

    auto_adjust_columns(ws2, max_cols=11)
    ws2.freeze_panes = "A6"

    # Sheet 3: Resumen Ejecutivo y Estadísticas
    ws3 = wb.create_sheet(title="Resumen Ejecutivo")
    apply_title_header(
        ws3,
        "MAPA CJB - RESUMEN EJECUTIVO Y ESTADÍSTICAS (CIUDAD JUAN BOSCH)",
        "Distribución Consolidada por Estatus de Obra y Tipo de Uso de Suelo",
        "Total: 213 Lotes"
    )

    # Tabla 1: Resumen por Estatus
    ws3.cell(row=6, column=1, value="DISTRIBUCIÓN POR ESTATUS DE OBRA").font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY)
    headers_st = ["Estatus", "Cantidad de Lotes", "% Lotes", "Superficie Total (m²)", "% Superficie"]
    for col_idx, h in enumerate(headers_st, 1):
        c = ws3.cell(row=7, column=col_idx, value=h)
        c.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
        c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = get_border()

    st_grouped = df_master.groupby("Estatus de Obra").agg(
        lotes=('Lote ID', 'count'),
        area=('Superficie (m²)', 'sum')
    ).reset_index().sort_values(by='lotes', ascending=False)

    tot_lotes_cjb = len(df_master)
    tot_area_cjb = df_master['Superficie (m²)'].sum()

    for idx, r in st_grouped.iterrows():
        r_row = 8 + idx
        ws3.row_dimensions[r_row].height = 20
        est = r["Estatus de Obra"]
        cnt = r["lotes"]
        ar = r["area"]

        c_e = ws3.cell(row=r_row, column=1, value=est)
        status_style = STATUS_COLORS.get(est, {'fill': 'F1F5F9', 'font': '475569'})
        c_e.fill = PatternFill(start_color=status_style['fill'], end_color=status_style['fill'], fill_type="solid")
        c_e.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=status_style['font'])
        c_e.alignment = Alignment(horizontal="center", vertical="center")
        c_e.border = border_thin

        c_cnt = ws3.cell(row=r_row, column=2, value=cnt)
        c_cnt.number_format = "#,##0"
        c_cnt.alignment = Alignment(horizontal="right", vertical="center")
        c_cnt.border = border_thin

        c_pct = ws3.cell(row=r_row, column=3, value=cnt / tot_lotes_cjb)
        c_pct.number_format = "0.0%"
        c_pct.alignment = Alignment(horizontal="right", vertical="center")
        c_pct.border = border_thin

        c_ar = ws3.cell(row=r_row, column=4, value=ar)
        c_ar.number_format = "#,##0.00"
        c_ar.alignment = Alignment(horizontal="right", vertical="center")
        c_ar.border = border_thin

        c_arpct = ws3.cell(row=r_row, column=5, value=ar / tot_area_cjb)
        c_arpct.number_format = "0.0%"
        c_arpct.alignment = Alignment(horizontal="right", vertical="center")
        c_arpct.border = border_thin

    st_total_row = 8 + len(st_grouped)
    ws3.cell(row=st_total_row, column=1, value="TOTAL:").font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=1).border = get_total_border()
    ws3.cell(row=st_total_row, column=2, value=tot_lotes_cjb).number_format = "#,##0"
    ws3.cell(row=st_total_row, column=2).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=2).border = get_total_border()
    ws3.cell(row=st_total_row, column=3, value=1.0).number_format = "0.0%"
    ws3.cell(row=st_total_row, column=3).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=3).border = get_total_border()
    ws3.cell(row=st_total_row, column=4, value=tot_area_cjb).number_format = "#,##0.00"
    ws3.cell(row=st_total_row, column=4).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=4).border = get_total_border()
    ws3.cell(row=st_total_row, column=5, value=1.0).number_format = "0.0%"
    ws3.cell(row=st_total_row, column=5).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=5).border = get_total_border()

    # Tabla 2: Resumen por Tipo de Suelo
    t_start = st_total_row + 3
    ws3.cell(row=t_start, column=1, value="DISTRIBUCIÓN POR TIPO DE USO DE SUELO").font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY)
    headers_tp = ["Tipo de Suelo", "Cantidad de Lotes", "% Lotes", "Superficie Total (m²)", "% Superficie"]
    for col_idx, h in enumerate(headers_tp, 1):
        c = ws3.cell(row=t_start+1, column=col_idx, value=h)
        c.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
        c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = get_border()

    tp_grouped = df_master.groupby("Tipo Suelo").agg(
        lotes=('Lote ID', 'count'),
        area=('Superficie (m²)', 'sum')
    ).reset_index().sort_values(by='lotes', ascending=False)

    for idx, r in tp_grouped.iterrows():
        r_row = t_start + 2 + idx
        ws3.row_dimensions[r_row].height = 20
        tp = r["Tipo Suelo"]
        cnt = r["lotes"]
        ar = r["area"]

        c_t = ws3.cell(row=r_row, column=1, value=tp)
        c_t.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_TEXT_MAIN)
        c_t.alignment = Alignment(horizontal="center", vertical="center")
        c_t.border = border_thin

        c_cnt = ws3.cell(row=r_row, column=2, value=cnt)
        c_cnt.number_format = "#,##0"
        c_cnt.alignment = Alignment(horizontal="right", vertical="center")
        c_cnt.border = border_thin

        c_pct = ws3.cell(row=r_row, column=3, value=cnt / tot_lotes_cjb)
        c_pct.number_format = "0.0%"
        c_pct.alignment = Alignment(horizontal="right", vertical="center")
        c_pct.border = border_thin

        c_ar = ws3.cell(row=r_row, column=4, value=ar)
        c_ar.number_format = "#,##0.00"
        c_ar.alignment = Alignment(horizontal="right", vertical="center")
        c_ar.border = border_thin

        c_arpct = ws3.cell(row=r_row, column=5, value=ar / tot_area_cjb)
        c_arpct.number_format = "0.0%"
        c_arpct.alignment = Alignment(horizontal="right", vertical="center")
        c_arpct.border = border_thin

    tp_total_row = t_start + 2 + len(tp_grouped)
    ws3.cell(row=tp_total_row, column=1, value="TOTAL:").font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=1).border = get_total_border()
    ws3.cell(row=tp_total_row, column=2, value=tot_lotes_cjb).number_format = "#,##0"
    ws3.cell(row=tp_total_row, column=2).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=2).border = get_total_border()
    ws3.cell(row=tp_total_row, column=3, value=1.0).number_format = "0.0%"
    ws3.cell(row=tp_total_row, column=3).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=3).border = get_total_border()
    ws3.cell(row=tp_total_row, column=4, value=tot_area_cjb).number_format = "#,##0.00"
    ws3.cell(row=tp_total_row, column=4).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=4).border = get_total_border()
    ws3.cell(row=tp_total_row, column=5, value=1.0).number_format = "0.0%"
    ws3.cell(row=tp_total_row, column=5).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=5).border = get_total_border()

    auto_adjust_columns(ws3, max_cols=5)

    wb.save(dest_path)
    print(f"[OK] CJB Excel saved successfully: {dest_path}")


def build_cpag_excel(dest_path):
    print("Building CPAG Excel...")
    csv_path = os.path.join(ROOT, "Datos de CPAG.csv")
    df_cpag = pd.read_csv(csv_path, encoding='utf-8')

    wb = openpyxl.Workbook()

    # Sheet 1: Listado de Lotes con las 4 columnas requeridas
    ws1 = wb.active
    ws1.title = "Listado de Lotes"

    apply_title_header(
        ws1,
        "MAPA CPAG - LISTADO OFICIAL DE LOTES (CIUDAD PDTE. ANTONIO GUZMÁN)",
        "Listado con ID, Nombre Comercial, Descripción y Estatus",
        "Universo: 98 Lotes (Master Plan Santiago)"
    )

    headers_ws1 = ["No.", "ID Lote", "Nombre Comercial", "Descripción", "Estatus"]
    ws1.row_dimensions[5].height = 26

    for col_idx, h in enumerate(headers_ws1, 1):
        cell = ws1.cell(row=5, column=col_idx, value=h)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_WHITE)
        cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = get_border()

    border_thin = get_border()
    row_start = 6

    # Mapeo descriptivo por tipo para CPAG
    desc_map = {
        'HABITACIONAL': 'Lote Habitacional (Vivienda Multifamiliar)',
        'COMERCIAL': 'Lote Comercial (Locales y Comercios)',
        'INSTITUCIONAL': 'Lote Institucional (Equipamiento Comunitario)',
        'INFRAESTRUCTURA': 'Lote de Infraestructura y Servicios',
        'AREA VERDE': 'Área Verde y Recreativa'
    }

    parsed_cpag_rows = []

    for i, (_, r) in enumerate(df_cpag.iterrows(), 1):
        curr_row = row_start + i - 1
        ws1.row_dimensions[curr_row].height = 20
        is_zebra = (i % 2 == 0)
        row_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_zebra else PatternFill(fill_type=None)

        lote_id = str(r["id"]).strip()
        tipo_raw = str(r["tipo"]).strip().upper() if pd.notnull(r["tipo"]) else ""
        subtipo_raw = str(r["subtipo"]).strip() if pd.notnull(r["subtipo"]) else ""
        nombre_raw = str(r["nombre"]).strip() if pd.notnull(r["nombre"]) else ""
        estado_raw = str(r["estado"]).strip() if pd.notnull(r["estado"]) else ""
        dev_raw = str(r["desarrollador"]).strip() if pd.notnull(r["desarrollador"]) else ""
        obs_raw = str(r["observaciones"]).strip() if pd.notnull(r["observaciones"]) else ""

        # Limpiar área
        area_str = str(r["area_m2"]).replace("m²", "").replace("m2", "").replace(",", "").strip()
        try:
            area_val = float(area_str)
        except:
            area_val = 0.0

        # Determinar valores pulidos
        if lote_id in ["03-03-01", "03-02-01"]:
            nombre_comercial = f"Proyecto Habitacional ({dev_raw})"
            desc = "Lote Habitacional - Multifamiliar (Adjudicado)"
            estatus = "Adjudicado"
            desarrollador = dev_raw
        else:
            nombre_comercial = "Por Definir"
            desc = desc_map.get(tipo_raw, f"Lote {tipo_raw.title()}")
            estatus = "No Iniciado"
            desarrollador = "Sin Desarrollador Asignado"

        parsed_cpag_rows.append({
            'no': i,
            'id': lote_id,
            'nombre_comercial': nombre_comercial,
            'desc': desc,
            'estatus': estatus,
            'tipo': tipo_raw.title() if tipo_raw != "AREA VERDE" else "Área Verde",
            'desarrollador': desarrollador,
            'area': area_val,
            'observaciones': obs_raw
        })

        # 1. No.
        c1 = ws1.cell(row=curr_row, column=1, value=i)
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_TEXT_MAIN)
        c1.fill = row_fill
        c1.border = border_thin

        # 2. ID Lote (Texto)
        c2 = ws1.cell(row=curr_row, column=2, value=lote_id)
        c2.number_format = "@"
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY)
        c2.fill = row_fill
        c2.border = border_thin

        # 3. Nombre Comercial
        c3 = ws1.cell(row=curr_row, column=3, value=nombre_comercial)
        c3.alignment = Alignment(horizontal="left", vertical="center")
        c3.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_TEXT_MAIN)
        c3.fill = row_fill
        c3.border = border_thin

        # 4. Descripción
        c4 = ws1.cell(row=curr_row, column=4, value=desc)
        c4.alignment = Alignment(horizontal="left", vertical="center")
        c4.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_TEXT_MAIN)
        c4.fill = row_fill
        c4.border = border_thin

        # 5. Estatus
        c5 = ws1.cell(row=curr_row, column=5, value=estatus)
        c5.alignment = Alignment(horizontal="center", vertical="center")
        status_style = STATUS_COLORS.get(estatus, {'fill': 'F1F5F9', 'font': '475569'})
        c5.fill = PatternFill(start_color=status_style['fill'], end_color=status_style['fill'], fill_type="solid")
        c5.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=status_style['font'])
        c5.border = border_thin

    # Fila de totales Sheet 1
    total_row = row_start + len(df_cpag)
    ws1.row_dimensions[total_row].height = 22
    ws1.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    tot_lbl = ws1.cell(row=total_row, column=1, value="TOTAL LOTES REGISTRADOS EN MAPA CPAG:")
    tot_lbl.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    tot_lbl.alignment = Alignment(horizontal="right", vertical="center")
    tot_lbl.border = get_total_border()

    tot_val = ws1.cell(row=total_row, column=5, value=len(df_cpag))
    tot_val.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    tot_val.alignment = Alignment(horizontal="center", vertical="center")
    tot_val.border = get_total_border()

    auto_adjust_columns(ws1, max_cols=5)
    ws1.freeze_panes = "A6"

    # Sheet 2: Detalle Técnico Completo CPAG
    ws2 = wb.create_sheet(title="Detalle Técnico Completo")
    apply_title_header(
        ws2,
        "MAPA CPAG - INVENTARIO TÉCNICO DETALLADO (CIUDAD PDTE. ANTONIO GUZMÁN)",
        "Base Técnica de Zonificación, Superficie, Desarrolladores y Observaciones",
        "Universo: 98 Lotes"
    )

    headers_ws2 = [
        "No.", "ID Lote", "Nombre Comercial", "Tipo de Suelo", "Descripción del Uso",
        "Estatus", "Desarrollador / Adjudicatario", "Superficie (m²)", "Observaciones / Situación Financiera"
    ]
    ws2.row_dimensions[5].height = 26
    for col_idx, h in enumerate(headers_ws2, 1):
        cell = ws2.cell(row=5, column=col_idx, value=h)
        cell.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
        cell.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = get_border()

    for idx, item in enumerate(parsed_cpag_rows, 1):
        curr_row = row_start + idx - 1
        ws2.row_dimensions[curr_row].height = 20
        is_zebra = (idx % 2 == 0)
        row_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid") if is_zebra else PatternFill(fill_type=None)

        ws2.cell(row=curr_row, column=1, value=item['no']).alignment = Alignment(horizontal="center", vertical="center")

        c_id = ws2.cell(row=curr_row, column=2, value=item['id'])
        c_id.number_format = "@"
        c_id.alignment = Alignment(horizontal="center", vertical="center")
        c_id.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_NAVY)

        ws2.cell(row=curr_row, column=3, value=item['nombre_comercial']).alignment = Alignment(horizontal="left", vertical="center")
        ws2.cell(row=curr_row, column=4, value=item['tipo']).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=curr_row, column=5, value=item['desc']).alignment = Alignment(horizontal="left", vertical="center")

        c_st = ws2.cell(row=curr_row, column=6, value=item['estatus'])
        c_st.alignment = Alignment(horizontal="center", vertical="center")
        status_style = STATUS_COLORS.get(item['estatus'], {'fill': 'F1F5F9', 'font': '475569'})
        c_st.fill = PatternFill(start_color=status_style['fill'], end_color=status_style['fill'], fill_type="solid")
        c_st.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=status_style['font'])

        ws2.cell(row=curr_row, column=7, value=item['desarrollador']).alignment = Alignment(horizontal="left", vertical="center")

        c_ar = ws2.cell(row=curr_row, column=8, value=item['area'])
        c_ar.number_format = "#,##0.00"
        c_ar.alignment = Alignment(horizontal="right", vertical="center")

        ws2.cell(row=curr_row, column=9, value=item['observaciones'] if item['observaciones'] else "-").alignment = Alignment(horizontal="left", vertical="center")

        for c_idx in range(1, 10):
            cell = ws2.cell(row=curr_row, column=c_idx)
            cell.border = border_thin
            if c_idx != 6:
                cell.fill = row_fill
                cell.font = Font(name=FONT_FAMILY, size=9.5, color=COLOR_TEXT_MAIN)

    # Totales hoja 2
    ws2.row_dimensions[total_row].height = 22
    ws2.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=7)
    tot_lbl2 = ws2.cell(row=total_row, column=1, value="SUPERFICIE TOTAL MAPA CPAG:")
    tot_lbl2.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    tot_lbl2.alignment = Alignment(horizontal="right", vertical="center")
    tot_lbl2.border = get_total_border()

    c_tot_ar2 = ws2.cell(row=total_row, column=8, value=f"=SUM(H6:H{total_row-1})")
    c_tot_ar2.number_format = "#,##0.00"
    c_tot_ar2.font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    c_tot_ar2.alignment = Alignment(horizontal="right", vertical="center")
    c_tot_ar2.border = get_total_border()

    c_blank = ws2.cell(row=total_row, column=9, value="")
    c_blank.border = get_total_border()

    auto_adjust_columns(ws2, max_cols=9)
    ws2.freeze_panes = "A6"

    # Sheet 3: Resumen Ejecutivo CPAG
    ws3 = wb.create_sheet(title="Resumen Ejecutivo")
    apply_title_header(
        ws3,
        "MAPA CPAG - RESUMEN EJECUTIVO Y ESTADÍSTICAS (CIUDAD PDTE. ANTONIO GUZMÁN)",
        "Distribución Consolidada por Estatus y Tipo de Uso de Suelo",
        "Total: 98 Lotes"
    )

    df_cpag_clean = pd.DataFrame(parsed_cpag_rows)
    tot_lotes_cpag = len(df_cpag_clean)
    tot_area_cpag = df_cpag_clean['area'].sum()

    # Tabla 1: Resumen por Estatus
    ws3.cell(row=6, column=1, value="DISTRIBUCIÓN POR ESTATUS DE LOTES").font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY)
    headers_st = ["Estatus", "Cantidad de Lotes", "% Lotes", "Superficie Total (m²)", "% Superficie"]
    for col_idx, h in enumerate(headers_st, 1):
        c = ws3.cell(row=7, column=col_idx, value=h)
        c.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
        c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = get_border()

    st_grouped = df_cpag_clean.groupby("estatus").agg(
        lotes=('id', 'count'),
        area=('area', 'sum')
    ).reset_index().sort_values(by='lotes', ascending=False)

    for idx, r in st_grouped.iterrows():
        r_row = 8 + idx
        ws3.row_dimensions[r_row].height = 20
        est = r["estatus"]
        cnt = r["lotes"]
        ar = r["area"]

        c_e = ws3.cell(row=r_row, column=1, value=est)
        status_style = STATUS_COLORS.get(est, {'fill': 'F1F5F9', 'font': '475569'})
        c_e.fill = PatternFill(start_color=status_style['fill'], end_color=status_style['fill'], fill_type="solid")
        c_e.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=status_style['font'])
        c_e.alignment = Alignment(horizontal="center", vertical="center")
        c_e.border = border_thin

        c_cnt = ws3.cell(row=r_row, column=2, value=cnt)
        c_cnt.number_format = "#,##0"
        c_cnt.alignment = Alignment(horizontal="right", vertical="center")
        c_cnt.border = border_thin

        c_pct = ws3.cell(row=r_row, column=3, value=cnt / tot_lotes_cpag)
        c_pct.number_format = "0.0%"
        c_pct.alignment = Alignment(horizontal="right", vertical="center")
        c_pct.border = border_thin

        c_ar = ws3.cell(row=r_row, column=4, value=ar)
        c_ar.number_format = "#,##0.00"
        c_ar.alignment = Alignment(horizontal="right", vertical="center")
        c_ar.border = border_thin

        c_arpct = ws3.cell(row=r_row, column=5, value=ar / tot_area_cpag)
        c_arpct.number_format = "0.0%"
        c_arpct.alignment = Alignment(horizontal="right", vertical="center")
        c_arpct.border = border_thin

    st_total_row = 8 + len(st_grouped)
    ws3.cell(row=st_total_row, column=1, value="TOTAL:").font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=1).border = get_total_border()
    ws3.cell(row=st_total_row, column=2, value=tot_lotes_cpag).number_format = "#,##0"
    ws3.cell(row=st_total_row, column=2).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=2).border = get_total_border()
    ws3.cell(row=st_total_row, column=3, value=1.0).number_format = "0.0%"
    ws3.cell(row=st_total_row, column=3).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=3).border = get_total_border()
    ws3.cell(row=st_total_row, column=4, value=tot_area_cpag).number_format = "#,##0.00"
    ws3.cell(row=st_total_row, column=4).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=4).border = get_total_border()
    ws3.cell(row=st_total_row, column=5, value=1.0).number_format = "0.0%"
    ws3.cell(row=st_total_row, column=5).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=st_total_row, column=5).border = get_total_border()

    # Tabla 2: Resumen por Tipo de Uso de Suelo
    t_start = st_total_row + 3
    ws3.cell(row=t_start, column=1, value="DISTRIBUCIÓN POR TIPO DE USO DE SUELO").font = Font(name=FONT_FAMILY, size=11, bold=True, color=COLOR_NAVY)
    headers_tp = ["Tipo de Suelo", "Cantidad de Lotes", "% Lotes", "Superficie Total (m²)", "% Superficie"]
    for col_idx, h in enumerate(headers_tp, 1):
        c = ws3.cell(row=t_start+1, column=col_idx, value=h)
        c.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_WHITE)
        c.fill = PatternFill(start_color=COLOR_NAVY, end_color=COLOR_NAVY, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = get_border()

    tp_grouped = df_cpag_clean.groupby("tipo").agg(
        lotes=('id', 'count'),
        area=('area', 'sum')
    ).reset_index().sort_values(by='lotes', ascending=False)

    for idx, r in tp_grouped.iterrows():
        r_row = t_start + 2 + idx
        ws3.row_dimensions[r_row].height = 20
        tp = r["tipo"]
        cnt = r["lotes"]
        ar = r["area"]

        c_t = ws3.cell(row=r_row, column=1, value=tp)
        c_t.font = Font(name=FONT_FAMILY, size=9.5, bold=True, color=COLOR_TEXT_MAIN)
        c_t.alignment = Alignment(horizontal="center", vertical="center")
        c_t.border = border_thin

        c_cnt = ws3.cell(row=r_row, column=2, value=cnt)
        c_cnt.number_format = "#,##0"
        c_cnt.alignment = Alignment(horizontal="right", vertical="center")
        c_cnt.border = border_thin

        c_pct = ws3.cell(row=r_row, column=3, value=cnt / tot_lotes_cpag)
        c_pct.number_format = "0.0%"
        c_pct.alignment = Alignment(horizontal="right", vertical="center")
        c_pct.border = border_thin

        c_ar = ws3.cell(row=r_row, column=4, value=ar)
        c_ar.number_format = "#,##0.00"
        c_ar.alignment = Alignment(horizontal="right", vertical="center")
        c_ar.border = border_thin

        c_arpct = ws3.cell(row=r_row, column=5, value=ar / tot_area_cpag)
        c_arpct.number_format = "0.0%"
        c_arpct.alignment = Alignment(horizontal="right", vertical="center")
        c_arpct.border = border_thin

    tp_total_row = t_start + 2 + len(tp_grouped)
    ws3.cell(row=tp_total_row, column=1, value="TOTAL:").font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=1).border = get_total_border()
    ws3.cell(row=tp_total_row, column=2, value=tot_lotes_cpag).number_format = "#,##0"
    ws3.cell(row=tp_total_row, column=2).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=2).border = get_total_border()
    ws3.cell(row=tp_total_row, column=3, value=1.0).number_format = "0.0%"
    ws3.cell(row=tp_total_row, column=3).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=3).border = get_total_border()
    ws3.cell(row=tp_total_row, column=4, value=tot_area_cpag).number_format = "#,##0.00"
    ws3.cell(row=tp_total_row, column=4).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=4).border = get_total_border()
    ws3.cell(row=tp_total_row, column=5, value=1.0).number_format = "0.0%"
    ws3.cell(row=tp_total_row, column=5).font = Font(name=FONT_FAMILY, size=10, bold=True, color=COLOR_NAVY)
    ws3.cell(row=tp_total_row, column=5).border = get_total_border()

    auto_adjust_columns(ws3, max_cols=5)

    wb.save(dest_path)
    print(f"[OK] CPAG Excel saved successfully: {dest_path}")

if __name__ == "__main__":
    out_cjb = os.path.join(ROOT, "LISTADO_LOTES_MAPA_CJB.xlsx")
    out_cpag = os.path.join(ROOT, "LISTADO_LOTES_MAPA_CPAG.xlsx")
    build_cjb_excel(out_cjb)
    build_cpag_excel(out_cpag)
